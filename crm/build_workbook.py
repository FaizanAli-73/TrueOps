"""Builds the Business Data Workbook.

    python build_workbook.py                    -> Excel version
    python build_workbook.py --gsheets          -> Google Sheets version (upload + convert in Drive)
    python build_workbook.py --out path.xlsx    -> custom output path

The Google Sheets version is the same workbook with the handful of Excel-only
constructs swapped for ones Sheets understands (static dropdown ranges, INDIRECT
in conditional formatting, no _xlfn prefixes) plus a WhatsApp Log tab for the bot.
"""
import os, re, sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import FormulaRule, ColorScaleRule
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter as GL

GSHEETS = "--gsheets" in sys.argv
BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "")
OUT = BASE + ("Business_Data_Workbook_GoogleSheets.xlsx" if GSHEETS else "Business_Data_Workbook.xlsx")
if "--out" in sys.argv:
    OUT = sys.argv[sys.argv.index("--out") + 1]
D = dt.date
T = dt.time

# ---------------------------------------------------------------- styles
FONT = "Arial"
NAVY, SLATE, GOLD, GREEN_T = "1F3864", "595959", "C9A227", "375623"
f_in = Font(name=FONT, size=10, color="0000FF")
f_auto = Font(name=FONT, size=10, color="000000")
f_warn = Font(name=FONT, size=10, color="C00000", bold=True)
f_hdr = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=16, bold=True, color=NAVY)
f_sub = Font(name=FONT, size=10, italic=True, color="595959")
f_bold = Font(name=FONT, size=10, bold=True)
f_body = Font(name=FONT, size=10)
f_sec = Font(name=FONT, size=12, bold=True, color="FFFFFF")
fill_in_h = PatternFill("solid", fgColor=NAVY)
fill_auto_h = PatternFill("solid", fgColor=SLATE)
fill_auto = PatternFill("solid", fgColor="F2F2F2")
fill_yellow = PatternFill("solid", fgColor="FFFF00")
fill_sec = PatternFill("solid", fgColor=NAVY)
fill_total = PatternFill("solid", fgColor="DDEBF7")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap_c = Alignment(horizontal="center", vertical="center", wrap_text=True)
wrap_l = Alignment(horizontal="left", vertical="top", wrap_text=True)
RED_F = PatternFill("solid", fgColor="F8CBAD")
AMB_F = PatternFill("solid", fgColor="FFE699")
GRN_F = PatternFill("solid", fgColor="C6EFCE")

TXT, DATE, TIME = "@", "mm/dd/yyyy", "h:mm AM/PM"
USD, USD0 = '$#,##0.00;($#,##0.00);"-"', '$#,##0;($#,##0);"-"'
PCT, HRS, INT, NUM = '0.0%;-0.0%;"-"', '0.00;-0.00;"-"', "0", "#,##0"
DAYS = '0;-0;0'
MULT = '0.0"x"'

HDR_ROW, FIRST = 4, 5

# ---------------------------------------------------------------- spec helpers
def I(key, header, width=13, fmt=None, dv=None, note=None, yellow=False):
    return dict(key=key, header=header, width=width, kind="in", fmt=fmt, dv=dv, note=note, yellow=yellow)

def A(key, header, f, width=13, fmt=None, note=None, guard=True, warn=False):
    return dict(key=key, header=header, width=width, kind="auto", fmt=fmt, f=f, note=note, guard=guard, warn=warn)

SHEETS = {}  # name -> dict(letters={key: letter})

def register(name, cols):
    SHEETS[name] = dict(letters={c["key"]: GL(i + 1) for i, c in enumerate(cols)}, cols=cols)

def W(sheet, key):
    L = SHEETS[sheet]["letters"][key]
    return f"'{sheet}'!${L}:${L}"

TOK = re.compile(r"\[(?:([A-Za-z &]+)!)?([A-Za-z0-9_]+)\]")
RNG = re.compile(r"\{rng:([A-Za-z0-9_]+)\}")

def render(tpl, sheet, r, first=None, last=None):
    lt = SHEETS[sheet]["letters"]
    def rep(m):
        sh, key = m.group(1), m.group(2)
        if sh:
            return W(sh, key)
        return f"{lt[key]}{r}"
    out = TOK.sub(rep, tpl)
    out = RNG.sub(lambda m: f"${lt[m.group(1)]}${first}:${lt[m.group(1)]}${last}", out)
    return out

YN = "List_YesNo"
PERIOD = 'IFERROR(IF(OR(AnalysisPeriod="All Time",YEAR({d})&""=AnalysisPeriod&""),1,0),0)'

def lookup(sheet, ret, key_here, key_there):
    return f'IFERROR(INDEX([{sheet}!{ret}],MATCH([{key_here}],[{sheet}!{key_there}],0)),"")'

# ================================================================ DATA SHEETS
clients = [
    I("ClientID", "Client ID", 10, TXT, note="Unique ID, e.g. C-0001. Never change or reuse an ID once used."),
    I("ClientType", "Client Type", 17, dv="List_ClientTypes"),
    I("FirstName", "First Name", 12), I("LastName", "Last Name", 13),
    I("Company", "Company (if a business)", 18),
    A("DisplayName", "Client Name", 'IF([Company]<>"",[Company],TRIM([FirstName]&" "&[LastName]))', 20),
    I("Phone", "Phone", 13, TXT), I("Email", "Email", 22),
    I("PrefContact", "Preferred Contact", 10, dv="List_Contact"),
    I("Language", "Preferred Language", 11),
    I("Address", "Billing Street Address", 22), I("City", "City", 13), I("State", "State", 6), I("ZIP", "ZIP", 7, TXT),
    I("LeadSource", "Original Lead Source", 20, dv="List_LeadSources", note="How they FIRST found you. Every quote also records its own lead source (a returning client is 'Repeat Client')."),
    I("ReferredBy", "Referred By (Client ID)", 11, TXT, note="If an existing client referred them, enter that client's ID. It counts toward the referrer's 'Referrals Made'."),
    I("FirstContact", "First Contact Date", 11, DATE),
    I("OptIn", "OK to Market To?", 9, dv=YN),
    I("PriceSens", "Price Sensitivity", 10, dv="List_Sensitivity", note="Your read on the client: how hard do they push on price?"),
    I("DNS", "Do Not Serve", 8, dv=YN, note="Flag clients you won't work with again (non-payment, abusive, etc.)."),
    I("Notes", "Notes", 28),
    A("Quotes", "# Quotes", 'COUNTIFS([Quotes!ClientID],[ClientID])', 8, INT),
    A("Won", "# Quotes Won", 'COUNTIFS([Quotes!ClientID],[ClientID],[Quotes!Status],"Won")', 8, INT),
    A("Completed", "# Completed Projects", 'COUNTIFS([Projects!ClientID],[ClientID],[Projects!Status],"Completed")', 10, INT),
    A("LTRev", "Lifetime Revenue $", 'SUMIFS([Projects!Revenue],[Projects!ClientID],[ClientID],[Projects!Status],"Completed")', 12, USD0),
    A("LTGP", "Lifetime Gross Profit $", 'SUMIFS([Projects!GrossProfit],[Projects!ClientID],[ClientID],[Projects!Status],"Completed")', 12, USD0),
    A("AvgGM", "Avg Gross Margin %", 'IF([LTRev]=0,"",[LTGP]/[LTRev])', 10, PCT),
    A("LastProject", "Last Project Completed", 'IF([Completed]=0,"",_xlfn.MAXIFS([Projects!ActualEnd],[Projects!ClientID],[ClientID],[Projects!Status],"Completed"))', 11, DATE),
    A("DaysSince", "Days Since Last Project", 'IF([LastProject]="","",TODAY()-[LastProject])', 10, INT, note="Use this to find past clients to check in with (e.g. 12+ months = time for a maintenance call)."),
    A("Repeat", "Repeat Client?", 'IF([Completed]>=2,"Yes","No")', 8),
    A("Referrals", "Referrals Made", 'COUNTIF([Clients!ReferredBy],[ClientID])', 9, INT),
    A("AvgSat", "Avg Satisfaction (1-10)", 'IFERROR(AVERAGEIFS([Projects!Satisfaction],[Projects!ClientID],[ClientID]),"")', 10, "0.0"),
    A("AvgDaysPay", "Avg Days to Pay", 'IFERROR(AVERAGEIFS([Invoices!DaysToPay],[Invoices!ClientID],[ClientID]),"")', 9, "0.0"),
    A("Tier", "Client Tier", 'IF([LTGP]>=TierA,"A",IF([LTGP]>=TierB,"B","C"))', 7, note="A / B / C by lifetime gross profit. Thresholds are on Settings. A-clients get first dibs on your schedule."),
    A("SortKey", "Sort Key (helper)", '[LTGP]+ROW()/1000000', 9, "0.00", note="Helper for the Top-10 clients list. Ignore."),
]
register("Clients", clients)

properties = [
    I("PropertyID", "Property ID", 10, TXT, note="Unique ID per job site, e.g. P-0001. One client can own many properties (landlords, property managers)."),
    I("ClientID", "Client ID (owner)", 10, TXT, dv="List_ClientIDs"),
    A("ClientName", "Client Name", lookup("Clients", "DisplayName", "ClientID", "ClientID"), 18),
    I("Street", "Street Address", 22), I("Unit", "Unit", 6), I("City", "City", 13), I("State", "State", 6), I("ZIP", "ZIP", 7, TXT),
    I("County", "County", 11), I("Area", "Neighborhood / Area", 14),
    I("DistanceMi", "Distance from Shop (mi, one way)", 10, "0.0"),
    I("DriveMin", "Drive Time (min, one way)", 9, INT),
    A("TravelZone", "Travel Zone", 'IF([DistanceMi]="","",IF([DistanceMi]<=Zone1Max,INDEX(List_Zones,1),IF([DistanceMi]<=Zone2Max,INDEX(List_Zones,2),IF([DistanceMi]<=Zone3Max,INDEX(List_Zones,3),INDEX(List_Zones,4)))))', 16),
    I("PropertyType", "Property Type", 16, dv="List_PropertyTypes"),
    I("Occupancy", "Occupancy", 14, dv="List_Occupancy"),
    I("YearBuilt", "Year Built", 8, INT),
    A("HomeAge", "Building Age (yrs)", 'IF([YearBuilt]="","",YEAR(TODAY())-[YearBuilt])', 8, INT),
    I("SqFt", "Square Feet", 9, NUM), I("Stories", "Stories", 7, INT), I("Beds", "Bedrooms", 7, INT), I("Baths", "Bathrooms", 7, "0.0"),
    I("Foundation", "Foundation", 15, dv="List_Foundation"),
    I("HVACType", "HVAC System Type", 18, dv="List_HVACTypes"),
    I("HVACYear", "HVAC Install Year", 8, INT),
    A("HVACAge", "HVAC Age (yrs)", 'IF([HVACYear]="","",YEAR(TODAY())-[HVACYear])', 7, INT),
    I("WHType", "Water Heater Type", 15, dv="List_WHTypes"),
    I("WHYear", "Water Heater Install Year", 8, INT),
    A("WHAge", "Water Heater Age (yrs)", 'IF([WHYear]="","",YEAR(TODAY())-[WHYear])', 7, INT),
    I("PanelAmps", "Electrical Panel (amps)", 8, INT),
    I("PanelBrand", "Panel Brand", 11),
    I("Pipe", "Supply Pipe Material", 14, dv="List_Pipe"),
    I("HOA", "HOA?", 6, dv=YN),
    I("Jurisdiction", "Permit Jurisdiction", 16),
    I("Access", "Access / Parking Notes", 22), I("Pets", "Pets", 10),
    I("Notes", "Notes", 24),
    A("ProjectsHere", "# Completed Projects Here", 'COUNTIFS([Projects!PropertyID],[PropertyID],[Projects!Status],"Completed")', 9, INT),
    A("RevenueHere", "Lifetime Revenue Here $", 'SUMIFS([Projects!Revenue],[Projects!PropertyID],[PropertyID],[Projects!Status],"Completed")', 11, USD0),
    A("Opportunity", "Future Work Flags", 'TRIM(IF(AND(ISNUMBER([HVACAge]),[HVACAge]>=HVACAgeFlag),"• HVAC replacement ","")&IF(AND(ISNUMBER([WHAge]),[WHAge]>=WHAgeFlag),"• Water heater ","")&IF(AND(ISNUMBER([PanelAmps]),[PanelAmps]<=PanelFlagAmps),"• Panel upgrade ","")&IF(OR([Pipe]="Galvanized Steel",[Pipe]="Polybutylene"),"• Repipe",""))', 26,
      note="Equipment you noticed on site that is near end of life = your next job. Age thresholds are on Settings. Call these clients before it fails."),
]
register("Properties", properties)

quotes = [
    I("QuoteID", "Quote ID", 9, TXT, note="Every estimate gets a row, won or lost. e.g. Q-0001."),
    I("DateReceived", "Lead Date", 11, DATE, note="Date the client first asked for this job."),
    I("ClientID", "Client ID", 9, TXT, dv="List_ClientIDs"),
    A("ClientName", "Client Name", lookup("Clients", "DisplayName", "ClientID", "ClientID"), 18),
    I("PropertyID", "Property ID", 9, TXT, dv="List_PropertyIDs"),
    I("LeadSource", "Lead Source (this job)", 20, dv="List_LeadSources"),
    I("ProjectType", "Project Type", 26, dv="List_ProjectTypes"),
    A("Trade", "Trade", 'IFERROR(INDEX(PT_Trades,MATCH([ProjectType],PT_Types,0)),"")', 13),
    I("Scope", "Scope Summary", 30), I("Urgency", "Urgency", 14, dv="List_Urgency"),
    I("SiteVisit", "Site Visit Date", 11, DATE), I("QuoteSent", "Quote Sent Date", 11, DATE),
    I("EstimatingHrs", "Hours Spent Estimating", 9, HRS, note="Site visit + drive + writing the quote. Tells you what lost quotes cost you."),
    I("EstLaborHrs", "Est. Labor Hours", 9, HRS, note="Total crew hours you expect (2 people x 8 hrs = 16)."),
    I("EstMaterials", "Est. Materials $", 11, USD), I("EstSubs", "Est. Subcontractor $", 11, USD),
    I("EstOther", "Est. Other Costs $ (permits, rental, disposal)", 12, USD),
    I("QuotedPrice", "Quoted Price $", 11, USD),
    I("CompetitorName", "Competitor (if known)", 16), I("CompetitorPrice", "Competitor's Price $", 11, USD, note="Ask! Lost-job feedback is the cheapest market research you'll ever get. Also log it on Market Rates."),
    I("Status", "Status", 11, dv="List_QuoteStatuses"),
    I("DecisionDate", "Decision Date", 11, DATE),
    I("LostReason", "Lost Reason", 18, dv="List_LostReasons"),
    I("ProjectID", "Project ID (if Won)", 9, TXT),
    I("FollowUps", "# Follow-Ups Made", 8, INT),
    I("Notes", "Notes", 24),
    A("DaysToQuote", "Days to Send Quote", 'IF(OR([DateReceived]="",[QuoteSent]=""),"",[QuoteSent]-[DateReceived])', 8, DAYS, note="Speed wins jobs. Track this and keep it low."),
    A("EstLaborCost", "Est. Labor Cost $", 'IF([EstLaborHrs]="","",[EstLaborHrs]*BlendedLaborRate)', 11, USD, note="Est. hours x your blended loaded labor rate (Pricing & Hiring tab)."),
    A("EstTotalCost", "Est. Total Cost $", 'N([EstLaborCost])+N([EstMaterials])+N([EstSubs])+N([EstOther])', 11, USD),
    A("EstGM", "Est. Gross Margin %", 'IF(N([QuotedPrice])=0,"",([QuotedPrice]-[EstTotalCost])/[QuotedPrice])', 9, PCT),
    A("EstGPHr", "Est. Gross Profit / Labor Hr", 'IF(OR(N([EstLaborHrs])=0,N([QuotedPrice])=0),"",([QuotedPrice]-[EstTotalCost])/[EstLaborHrs])', 10, USD),
    A("MinPrice", "Price Needed for Target Margin $", 'IF([EstTotalCost]=0,"",[EstTotalCost]/(1-TargetGM))', 12, USD0, note="Est. total cost ÷ (1 − target gross margin from Settings). Quote at or above this."),
    A("PriceCheck", "Price Check", 'IF([EstGM]="","",IF([EstGM]>=TargetGM,"✔ Meets target","⚠ Below target margin"))', 15),
    A("DaysToDecision", "Days Quote → Decision", 'IF(OR([DecisionDate]="",[QuoteSent]=""),"",[DecisionDate]-[QuoteSent])', 9, DAYS),
    A("DaysWaiting", "Days Waiting (open quotes)", 'IF(AND(OR([Status]="Sent",[Status]="Follow-Up"),[QuoteSent]<>""),TODAY()-[QuoteSent],"")', 9, INT, note="Open quotes: follow up at 2, 7 and 14 days."),
    A("InPeriod", "In Period? (helper)", PERIOD.format(d="[DateReceived]"), 8, INT),
]
register("Quotes", quotes)

projects = [
    I("ProjectID", "Project ID", 9, TXT, note="Unique job ID, e.g. J-0001. Use the SAME ID on Time Log, Expenses, Invoices, Change Orders."),
    I("QuoteID", "Quote ID", 9, TXT, note="Links to the estimate so you can compare estimated vs actual."),
    I("ClientID", "Client ID", 9, TXT, dv="List_ClientIDs"),
    A("ClientName", "Client Name", lookup("Clients", "DisplayName", "ClientID", "ClientID"), 18),
    I("PropertyID", "Property ID", 9, TXT, dv="List_PropertyIDs"),
    A("City", "City", lookup("Properties", "City", "PropertyID", "PropertyID"), 12),
    I("ProjectName", "Project Name", 24),
    I("ProjectType", "Project Type", 26, dv="List_ProjectTypes", note="Pick the closest type. This is what the analysis groups by, so be consistent. Add new types on Settings."),
    A("Trade", "Trade", 'IFERROR(INDEX(PT_Trades,MATCH([ProjectType],PT_Types,0)),"")', 13),
    I("Status", "Status", 11, dv="List_ProjectStatuses", note="Only COMPLETED projects (with an Actual End date) count in the analysis tabs."),
    I("ContractType", "Contract Type", 14, dv="List_ContractTypes"),
    I("Complexity", "Complexity (1-5)", 8, INT, dv=("whole", 1, 5), note="1 = simple / routine, 5 = hard / lots of unknowns. Lets you see if hard jobs are priced right."),
    I("Permit", "Permit Required?", 8, dv=YN),
    I("LeadTech", "Lead Worker ID", 8, TXT, dv="List_WorkerIDs"),
    I("SchedStart", "Scheduled Start", 11, DATE), I("SchedEnd", "Scheduled Finish", 11, DATE),
    I("ActualStart", "Actual Start", 11, DATE), I("ActualEnd", "Actual Finish", 11, DATE),
    I("ContractValue", "Original Contract $", 12, USD, note="Agreed price before change orders. For Time & Materials jobs, enter the final billed amount (excluding change orders)."),
    I("Inspection", "Passed Inspection 1st Try?", 9, dv="List_YesNoNA"),
    I("Satisfaction", "Client Satisfaction (1-10)", 9, INT, dv=("whole", 1, 10), note="Ask: 'On a scale of 1-10, how likely are you to recommend us?'"),
    I("Testimonial", "Testimonial Collected?", 9, dv=YN),
    I("Review", "Online Review Left?", 9, dv=YN),
    I("Photos", "Photos Folder Link", 16),
    I("Notes", "Notes", 26),
    A("ClientType", "Client Type", lookup("Clients", "ClientType", "ClientID", "ClientID"), 15),
    A("TravelZone", "Travel Zone", lookup("Properties", "TravelZone", "PropertyID", "PropertyID"), 15),
    A("LeadSource", "Lead Source", 'IFERROR(INDEX([Quotes!LeadSource],MATCH([QuoteID],[Quotes!QuoteID],0)),IFERROR(INDEX([Clients!LeadSource],MATCH([ClientID],[Clients!ClientID],0)),""))', 18, note="From the quote; falls back to the client's original lead source."),
    A("CalendarDays", "Calendar Days", 'IF(OR([ActualStart]="",[ActualEnd]=""),"",[ActualEnd]-[ActualStart]+1)', 8, INT),
    A("DaysLate", "Days Late (+) / Early (−)", 'IF(OR([SchedEnd]="",[ActualEnd]=""),"",[ActualEnd]-[SchedEnd])', 8, DAYS),
    A("OnTime", "On Time?", 'IF([DaysLate]="","",IF([DaysLate]<=0,"Yes","No"))', 7),
    A("COValue", "Approved Change Orders $", "SUMIFS([Change Orders!Price],[Change Orders!ProjectID],[ProjectID],[Change Orders!Status],\"Approved\")", 11, USD),
    A("Revenue", "Total Revenue $", 'N([ContractValue])+[COValue]', 12, USD),
    A("EstHours", "Est. Labor Hours (quote + COs)", 'IFERROR(N(INDEX([Quotes!EstLaborHrs],MATCH([QuoteID],[Quotes!QuoteID],0))),0)+SUMIFS([Change Orders!AddedHours],[Change Orders!ProjectID],[ProjectID],[Change Orders!Status],"Approved")', 10, HRS),
    A("EstCost", "Est. Direct Cost $ (quote + COs)", 'IFERROR(N(INDEX([Quotes!EstTotalCost],MATCH([QuoteID],[Quotes!QuoteID],0))),0)+SUMIFS([Change Orders!AddedCost],[Change Orders!ProjectID],[ProjectID],[Change Orders!Status],"Approved")+SUMIFS([Change Orders!AddedHours],[Change Orders!ProjectID],[ProjectID],[Change Orders!Status],"Approved")*BlendedLaborRate', 11, USD),
    A("LaborHours", "Actual Labor Hours", 'SUMIFS([Time Log!Hours],[Time Log!ProjectID],[ProjectID])', 9, HRS, note="All crew hours logged to this job on Time Log (including travel & material runs)."),
    A("ProductiveHours", "Wrench-Time Hours", 'SUMIFS([Time Log!Hours],[Time Log!ProjectID],[ProjectID],[Time Log!Productive],"Yes")', 9, HRS, note="Hours spent actually doing the work (demo, install, finish...). Excludes travel, pickups, waiting, callbacks."),
    A("NonProdHours", "Travel / Pickup / Waiting Hours", '[LaborHours]-[ProductiveHours]', 9, HRS),
    A("HoursVar", "Hours vs Estimate %", 'IF(OR([EstHours]=0,[LaborHours]=0),"",[LaborHours]/[EstHours]-1)', 9, PCT, note="+20% = took 20% longer than estimated. Consistently positive for a project type = you're under-estimating it."),
    A("LaborCost", "Labor Cost $", 'SUMIFS([Time Log!LaborCost],[Time Log!ProjectID],[ProjectID])', 11, USD),
    A("Materials", "Materials $", 'SUMIFS([Expenses!Total],[Expenses!ProjectID],[ProjectID],[Expenses!Category],"Materials")', 11, USD),
    A("Subs", "Subcontractors $", 'SUMIFS([Expenses!Total],[Expenses!ProjectID],[ProjectID],[Expenses!Category],"Subcontractor")', 11, USD),
    A("Equipment", "Equipment Rental $", 'SUMIFS([Expenses!Total],[Expenses!ProjectID],[ProjectID],[Expenses!Category],"Equipment Rental")', 10, USD),
    A("Permits", "Permits & Inspections $", 'SUMIFS([Expenses!Total],[Expenses!ProjectID],[ProjectID],[Expenses!Category],"Permits & Inspections")', 10, USD),
    A("Disposal", "Disposal $", 'SUMIFS([Expenses!Total],[Expenses!ProjectID],[ProjectID],[Expenses!Category],"Disposal / Dumpster")', 10, USD),
    A("OtherJob", "Other Job Costs $ (incl. card fees)", 'SUMIFS([Expenses!Total],[Expenses!ProjectID],[ProjectID],[Expenses!CostType],"Job Cost")-[Materials]-[Subs]-[Equipment]-[Permits]-[Disposal]+SUMIFS([Invoices!Fee],[Invoices!ProjectID],[ProjectID])', 10, USD),
    A("DirectCost", "Total Direct Cost $", '[LaborCost]+[Materials]+[Subs]+[Equipment]+[Permits]+[Disposal]+[OtherJob]', 12, USD),
    A("CostVar", "Cost vs Estimate %", 'IF(OR([EstCost]=0,[DirectCost]=0),"",[DirectCost]/[EstCost]-1)', 9, PCT),
    A("GrossProfit", "Gross Profit $", '[Revenue]-[DirectCost]', 12, USD),
    A("GM", "Gross Margin %", 'IF([Revenue]=0,"",[GrossProfit]/[Revenue])', 9, PCT),
    A("RevPerHr", "Revenue / Labor Hr", 'IF([LaborHours]=0,"",[Revenue]/[LaborHours])', 10, USD),
    A("GPPerHr", "Gross Profit / Labor Hr", 'IF([LaborHours]=0,"",[GrossProfit]/[LaborHours])', 10, USD, note="THE key number for choosing jobs: how much profit each hour of your crew's time produced on this job."),
    A("GPPerDay", "Gross Profit / Calendar Day", 'IF(N([CalendarDays])=0,"",[GrossProfit]/[CalendarDays])', 10, USD),
    A("OverheadAlloc", "Overhead Share $", '[LaborHours]*OverheadPerHour', 10, USD, note="Labor hours x overhead per project hour (Pricing & Hiring tab). Your share of insurance, truck, phone, etc."),
    A("NetProfit", "Net Profit $ (after overhead)", '[GrossProfit]-[OverheadAlloc]', 12, USD),
    A("NetMargin", "Net Margin %", 'IF([Revenue]=0,"",[NetProfit]/[Revenue])', 9, PCT),
    A("Invoiced", "Invoiced $", 'SUMIFS([Invoices!Amount],[Invoices!ProjectID],[ProjectID])', 11, USD),
    A("Collected", "Collected $", 'SUMIFS([Invoices!AmountPaid],[Invoices!ProjectID],[ProjectID])', 11, USD),
    A("Balance", "Balance Owed $", '[Revenue]-[Collected]', 11, USD),
    A("LastPayment", "Last Payment Date", 'IF([Collected]=0,"",_xlfn.MAXIFS([Invoices!DatePaid],[Invoices!ProjectID],[ProjectID]))', 11, DATE),
    A("DaysToCollect", "Days Finish → Paid in Full", 'IF(OR([LastPayment]="",[ActualEnd]="",[Balance]>0.004),"",[LastPayment]-[ActualEnd])', 9, DAYS),
    A("Callbacks", "# Callbacks", 'COUNTIFS([Callbacks!ProjectID],[ProjectID])', 8, INT),
    A("CallbackCost", "Callback Cost $", 'SUMIFS([Callbacks!TotalCost],[Callbacks!ProjectID],[ProjectID])', 10, USD),
    A("InPeriod", "In Analysis? (helper)", 'IFERROR(IF(AND([Status]="Completed",[ActualEnd]<>"",OR(AnalysisPeriod="All Time",YEAR([ActualEnd])&""=AnalysisPeriod&"")),1,0),0)', 8, INT, note="1 = completed and inside the Analysis Period chosen on the Dashboard."),
]
register("Projects", projects)

change_orders = [
    I("COID", "Change Order ID", 10, TXT), I("ProjectID", "Project ID", 9, TXT, dv="List_ProjectIDs"),
    A("ProjectName", "Project Name", lookup("Projects", "ProjectName", "ProjectID", "ProjectID"), 22),
    I("Date", "Date", 11, DATE), I("Description", "Description", 30),
    I("Reason", "Reason", 22, dv="List_COReasons", note="'Our Error (No Charge)' tracks rework you ate. Price it at $0 so you can see what mistakes cost."),
    I("AddedHours", "Added Labor Hours (est.)", 9, HRS), I("AddedCost", "Added Materials/Other Cost $ (est.)", 11, USD),
    I("Price", "Change Order Price $", 11, USD), I("Status", "Status", 10, dv="List_COStatuses"),
    I("ApprovedDate", "Approved Date", 11, DATE), I("Signed", "Signed by Client?", 9, dv=YN),
    I("Notes", "Notes", 24),
    A("COMargin", "Est. CO Margin %", 'IF(N([Price])=0,"",([Price]-N([AddedCost])-N([AddedHours])*BlendedLaborRate)/[Price])', 9, PCT),
]
register("Change Orders", change_orders)

time_log = [
    I("Date", "Date", 11, DATE, note="One row per person, per job, per task, per day. Log EVERY paid hour — including owners, travel and material runs."),
    I("WorkerID", "Worker ID", 8, TXT, dv="List_WorkerIDs"),
    A("WorkerName", "Worker", lookup("Team", "Name", "WorkerID", "WorkerID"), 14),
    I("ProjectID", "Project ID (or OVERHEAD)", 11, TXT, dv=("free", "List_ProjectIDs"), note="Job ID from Projects, or type OVERHEAD for time not tied to a job (admin, shop, training, estimates for jobs you didn't win)."),
    A("ProjectTrade", "Trade", 'IF([ProjectID]="OVERHEAD","Overhead",' + lookup("Projects", "Trade", "ProjectID", "ProjectID") + ')', 12),
    I("Task", "Task / Phase", 22, dv="List_Tasks", note="Which phase of the work. This is how you find where you're slow (see Time Analysis)."),
    A("Productive", "Wrench Time?", 'IFERROR(INDEX(Task_Productive,MATCH([Task],Task_Names,0)),"")', 8),
    I("Start", "Start Time", 9, TIME), I("End", "End Time", 9, TIME),
    I("BreakMin", "Unpaid Break (min)", 8, INT),
    I("ManualHours", "Hours (if no start/end)", 9, HRS, note="Only if you didn't record start/end times."),
    A("Hours", "Hours", 'IF(AND([Start]<>"",[End]<>""),ROUND(MOD([End]-[Start],1)*24-N([BreakMin])/60,2),N([ManualHours]))', 8, HRS),
    A("CostRate", "Loaded Cost Rate $/hr", 'IFERROR(INDEX([Team!LoadedRate],MATCH([WorkerID],[Team!WorkerID],0)),0)', 9, USD),
    A("LaborCost", "Labor Cost $", '[Hours]*[CostRate]', 10, USD),
    I("Notes", "Notes", 26),
    A("RevEarned", "Revenue Earned $ (attributed)", 'IFERROR(IF(INDEX([Projects!Status],MATCH([ProjectID],[Projects!ProjectID],0))="Completed",[Hours]*INDEX([Projects!RevPerHr],MATCH([ProjectID],[Projects!ProjectID],0)),0),0)', 11, USD, note="This person's share of the job's revenue, by hours worked. Only counts completed jobs."),
    A("GPEarned", "Gross Profit Earned $ (attributed)", 'IFERROR(IF(INDEX([Projects!Status],MATCH([ProjectID],[Projects!ProjectID],0))="Completed",[Hours]*INDEX([Projects!GPPerHr],MATCH([ProjectID],[Projects!ProjectID],0)),0),0)', 11, USD),
    A("InPeriod", "In Period? (helper)", PERIOD.format(d="[Date]"), 8, INT),
    A("Check", "Data Check", 'IF([WorkerID]="","⚠ Worker ID missing",IF([ProjectID]="","⚠ Project ID missing (or OVERHEAD)",IF(AND([ProjectID]<>"OVERHEAD",ISNA(MATCH([ProjectID],[Projects!ProjectID],0))),"⚠ Unknown Project ID",IF([Task]="","⚠ Task missing",IF([Hours]<=0,"⚠ Hours are zero",IF([Hours]>14,"⚠ Over 14 hrs — check",""))))))', 22, warn=True),
]
register("Time Log", time_log)

expenses = [
    I("Date", "Date", 11, DATE, note="Every receipt / bill gets a row."),
    I("ProjectID", "Project ID (or OVERHEAD)", 11, TXT, dv=("free", "List_ProjectIDs"), note="Job ID if the cost was for a specific job; OVERHEAD for general business costs (insurance, truck payment, tools, marketing...)."),
    I("Category", "Category", 24, dv="List_ExpCategories"),
    A("CostType", "Job Cost or Overhead", 'IFERROR(INDEX(Cat_Types,MATCH([Category],Cat_Names,0)),"")', 11),
    I("Vendor", "Vendor", 18, dv=("free", "List_Vendors")),
    I("Description", "Description", 28),
    I("Amount", "Amount (before tax) $", 11, USD), I("Tax", "Sales Tax $", 9, USD),
    A("Total", "Total $", 'N([Amount])+N([Tax])', 11, USD),
    I("PaidWith", "Paid With", 16, dv="List_PaidWith"),
    I("Channel", "Marketing Channel (marketing spend only)", 20, dv="List_LeadSources", note="For Marketing & Advertising only: which lead source this money was spent on. Powers cost-per-lead on the Marketing tab."),
    I("Receipt", "Receipt Link / #", 14), I("Notes", "Notes", 22),
    A("InPeriod", "In Period? (helper)", PERIOD.format(d="[Date]"), 8, INT),
    A("Check", "Data Check", 'IF([ProjectID]="","⚠ Enter Project ID or OVERHEAD",IF([Category]="","⚠ Category missing",IF(AND([CostType]="Job Cost",[ProjectID]="OVERHEAD"),"⚠ Job cost needs a Project ID",IF(AND([CostType]="Overhead",[ProjectID]<>"OVERHEAD"),"⚠ Overhead category on a job",IF(AND([ProjectID]<>"OVERHEAD",ISNA(MATCH([ProjectID],[Projects!ProjectID],0))),"⚠ Unknown Project ID",IF(AND([Category]="Marketing & Advertising",[Channel]=""),"Tip: add marketing channel",""))))))', 24, warn=True),
]
register("Expenses", expenses)

invoices = [
    I("InvoiceNo", "Invoice #", 9, TXT), I("ProjectID", "Project ID", 9, TXT, dv="List_ProjectIDs"),
    A("ClientID", "Client ID", lookup("Projects", "ClientID", "ProjectID", "ProjectID"), 9),
    A("ClientName", "Client Name", lookup("Projects", "ClientName", "ProjectID", "ProjectID"), 18),
    I("InvoiceType", "Invoice Type", 12, dv="List_InvoiceTypes"),
    I("InvoiceDate", "Invoice Date", 11, DATE),
    I("Terms", "Terms (days)", 7, INT, note="Days until due. Leave blank to use the default on Settings."),
    I("Amount", "Amount Invoiced $", 11, USD), I("AmountPaid", "Amount Paid $", 11, USD),
    I("DatePaid", "Date Paid", 11, DATE), I("Method", "Payment Method", 15, dv="List_PaymentMethods"),
    I("Fee", "Card / Processing Fee $", 9, USD, note="Fees taken out of the payment (card, financing). Counted as a job cost on the project — do NOT also log them on Expenses."),
    I("Notes", "Notes", 22),
    A("DueDate", "Due Date", 'IF([InvoiceDate]="","",[InvoiceDate]+IF([Terms]="",DefaultTerms,[Terms]))', 11, DATE),
    A("Balance", "Balance $", 'N([Amount])-N([AmountPaid])', 11, USD),
    A("Status", "Status", 'IF(N([Amount])=0,"",IF([Balance]<=0.004,"Paid",IF(AND([DueDate]<>"",TODAY()>[DueDate]),"OVERDUE",IF(N([AmountPaid])>0,"Partial","Open"))))', 9),
    A("DaysToPay", "Days to Pay", 'IF(OR([DatePaid]="",[InvoiceDate]=""),"",[DatePaid]-[InvoiceDate])', 7, DAYS),
    A("DaysOverdue", "Days Overdue", 'IF([Status]="OVERDUE",TODAY()-[DueDate],"")', 7, INT),
    A("PaidInPeriod", "Paid In Period? (helper)", PERIOD.format(d="[DatePaid]"), 8, INT),
]
register("Invoices", invoices)

team = [
    I("WorkerID", "Worker ID", 8, TXT, note="e.g. W-01. Owners get a row too."),
    I("Name", "Name", 16), I("Role", "Role", 14, dv="List_Roles"),
    I("EmpType", "Employment Type", 13, dv="List_EmpTypes"),
    I("PrimaryTrade", "Primary Trade", 13, dv="List_Trades"),
    I("OtherTrades", "Other Trades / Skills", 18),
    I("Skill", "Skill Level", 11, dv="List_Skill"),
    I("Licenses", "Licenses & Certifications", 20, note="e.g. EPA 608 Universal, Journeyman Electrician #12345, OSHA 10"),
    I("LicenseExpiry", "Next License Expiry", 11, DATE),
    I("StartDate", "Start Date", 11, DATE),
    I("Status", "Status", 8, dv="List_WorkerStatus"),
    I("BasePay", "Base Pay $/hr", 9, USD, note="What you pay per hour. OWNERS: enter what you'd have to pay someone to replace you on the tools. Leaving owners at $0 makes every job look more profitable than it is."),
    I("BurdenPct", "Burden % (blank = default)", 9, PCT, note="Payroll taxes + workers' comp + benefits as % of pay. Blank uses the Settings default (0% for 1099 contractors)."),
    I("Phone", "Phone" + (" (WhatsApp)" if GSHEETS else ""), 13, TXT,
      note=("The number this person texts the WhatsApp bot from, with country code (e.g. 13125550142). Only numbers listed here can write to this workbook." if GSHEETS else None)),
    I("Email", "Email", 18), I("Emergency", "Emergency Contact", 16),
    I("Notes", "Notes", 20),
    A("LicenseFlag", "License Status", 'IF([LicenseExpiry]="","",IF([LicenseExpiry]<TODAY(),"EXPIRED",IF([LicenseExpiry]-TODAY()<=60,"Renew soon","OK")))', 9),
    A("BurdenUsed", "Burden % Used", 'IF([BurdenPct]="",IF([EmpType]="1099 Contractor",0,BurdenDefault),[BurdenPct])', 8, PCT),
    A("LoadedRate", "Loaded Cost $/hr", 'N([BasePay])*(1+[BurdenUsed])', 9, USD, note="What one hour of this person really costs you."),
    A("Hours", "Hours Logged (period)", 'SUMIFS([Time Log!Hours],[Time Log!WorkerID],[WorkerID],[Time Log!InPeriod],1)', 9, HRS),
    A("ProjHours", "Hours on Jobs (period)", 'SUMIFS([Time Log!Hours],[Time Log!WorkerID],[WorkerID],[Time Log!InPeriod],1,[Time Log!ProjectID],"<>OVERHEAD")', 9, HRS),
    A("WrenchHours", "Wrench-Time Hours (period)", 'SUMIFS([Time Log!Hours],[Time Log!WorkerID],[WorkerID],[Time Log!InPeriod],1,[Time Log!Productive],"Yes")', 9, HRS),
    A("Utilization", "Wrench-Time %", 'IF([Hours]=0,"",[WrenchHours]/[Hours])', 8, PCT, note="Share of paid hours spent actually doing the work. 70-80%+ is strong for a field tech."),
    A("LaborCostP", "Labor Cost $ (period)", 'SUMIFS([Time Log!LaborCost],[Time Log!WorkerID],[WorkerID],[Time Log!InPeriod],1)', 10, USD),
    A("RevGen", "Revenue Generated $ (period)", 'SUMIFS([Time Log!RevEarned],[Time Log!WorkerID],[WorkerID],[Time Log!InPeriod],1)', 11, USD0),
    A("GPGen", "Gross Profit Generated $ (period)", 'SUMIFS([Time Log!GPEarned],[Time Log!WorkerID],[WorkerID],[Time Log!InPeriod],1)', 11, USD0),
    A("RevPerHr", "Revenue / Hour Paid", 'IF([Hours]=0,"",[RevGen]/[Hours])', 9, USD),
    A("ValueMult", "Revenue per $1 of Labor Cost", 'IF([LaborCostP]=0,"",[RevGen]/[LaborCostP])', 9, MULT, note="How many dollars of revenue this person's hours bring in for every $1 they cost. Compare across the team."),
    A("ProjectsLed", "Jobs Led (period)", 'COUNTIFS([Projects!LeadTech],[WorkerID],[Projects!InPeriod],1)', 7, INT),
    A("SatLed", "Avg Satisfaction on Jobs Led", 'IFERROR(AVERAGEIFS([Projects!Satisfaction],[Projects!LeadTech],[WorkerID],[Projects!InPeriod],1),"")', 9, "0.0"),
    A("CallbacksW", "Callbacks Caused (all time)", 'COUNTIFS([Callbacks!Worker],[WorkerID],[Callbacks!RootCause],"Workmanship")+COUNTIFS([Callbacks!Worker],[WorkerID],[Callbacks!RootCause],"Installation Oversight")', 9, INT),
]
register("Team", team)

vendors = [
    I("VendorID", "Vendor ID", 8, TXT), I("VendorName", "Vendor Name", 20, note="Use this exact name in the Vendor column on Expenses."),
    I("VendorType", "Vendor Type", 15, dv="List_VendorTypes"),
    I("Specialty", "Trade / Specialty", 13, dv="List_Trades"),
    I("Contact", "Contact Person", 14), I("Phone", "Phone", 12, TXT), I("Email", "Email", 18),
    I("AccountNo", "Account #", 10, TXT), I("Terms", "Payment Terms", 9, note="e.g. Net 30, COD"),
    I("Discount", "Contractor Discount %", 8, PCT),
    I("W9", "W-9 on File?", 7, dv=YN), I("COI", "Insurance Cert on File?", 8, dv=YN, note="For subcontractors: always get their certificate of insurance before they work your job."),
    I("COIExpiry", "Insurance Expiry", 11, DATE), I("LicenseNo", "License #", 11, TXT),
    I("Quality", "Quality (1-5)", 7, INT, dv=("whole", 1, 5)), I("Reliability", "Reliability (1-5)", 7, INT, dv=("whole", 1, 5)),
    I("Notes", "Notes", 22),
    A("COIFlag", "Insurance Status", 'IF([COIExpiry]="","",IF([COIExpiry]<TODAY(),"EXPIRED",IF([COIExpiry]-TODAY()<=30,"Renew soon","OK")))', 9),
    A("Spend", "Spend $ (period)", 'SUMIFS([Expenses!Total],[Expenses!Vendor],[VendorName],[Expenses!InPeriod],1)', 11, USD0),
    A("Purchases", "# Purchases (period)", 'COUNTIFS([Expenses!Vendor],[VendorName],[Expenses!InPeriod],1)', 8, INT),
    A("LastPurchase", "Last Purchase", 'IF(COUNTIFS([Expenses!Vendor],[VendorName])=0,"",_xlfn.MAXIFS([Expenses!Date],[Expenses!Vendor],[VendorName]))', 11, DATE),
]
register("Vendors", vendors)

callbacks = [
    I("CallbackID", "Callback ID", 9, TXT, note="Any time you go back to fix something on a finished job."),
    I("ProjectID", "Original Project ID", 9, TXT, dv="List_ProjectIDs"),
    A("ProjectName", "Project Name", lookup("Projects", "ProjectName", "ProjectID", "ProjectID"), 20),
    A("Trade", "Trade", lookup("Projects", "Trade", "ProjectID", "ProjectID"), 12),
    I("DateReported", "Date Reported", 11, DATE), I("Issue", "Issue", 28),
    I("RootCause", "Root Cause", 18, dv="List_CallbackCauses"),
    I("Worker", "Responsible Worker ID", 9, TXT, dv="List_WorkerIDs"),
    I("HoursSpent", "Hours Spent Fixing", 8, HRS, note="Also log this time on Time Log against the ORIGINAL project with task 'Warranty / Callback' so it hits that job's profit."),
    I("MaterialCost", "Material Cost $", 10, USD, note="Also log the receipt on Expenses against the ORIGINAL project."),
    I("DateResolved", "Date Resolved", 11, DATE), I("NoCharge", "No Charge to Client?", 8, dv=YN),
    I("Notes", "Notes / Lesson Learned", 26),
    A("TotalCost", "Est. Callback Cost $", 'N([HoursSpent])*BlendedLaborRate+N([MaterialCost])', 10, USD),
    A("DaysToResolve", "Days to Resolve", 'IF(OR([DateResolved]="",[DateReported]=""),"",[DateResolved]-[DateReported])', 8, DAYS),
    A("DaysAfter", "Days After Job Finished", 'IFERROR(IF(INDEX([Projects!ActualEnd],MATCH([ProjectID],[Projects!ProjectID],0))="","",[DateReported]-INDEX([Projects!ActualEnd],MATCH([ProjectID],[Projects!ProjectID],0))),"")', 8, DAYS),
]
register("Callbacks", callbacks)

market = [
    I("Date", "Date Observed", 11, DATE, note="Every competitor price you learn about. The more rows, the better your pricing decisions."),
    I("Competitor", "Competitor", 18), I("Trade", "Trade", 13, dv="List_Trades"),
    I("ProjectType", "Project Type (if known)", 24, dv="List_ProjectTypes"),
    I("Basis", "Pricing Basis", 16, dv="List_PricingBasis"),
    I("Price", "Rate or Price $", 10, USD, note="Hourly rate, whole-job price, or price per sq ft / per unit."),
    I("Qty", "Quantity (sq ft, units, hrs; blank = 1)", 9, NUM),
    I("OurHours", "Our Est. Hours for Same Job (optional)", 9, HRS, note="Lets the workbook turn a flat price into an effective $/hr. If blank, uses your average hours for that project type."),
    I("OurPrice", "Our Price for Same Job $ (optional)", 10, USD),
    I("Source", "Source", 22, dv="List_MarketSources"), I("AreaZip", "Area / ZIP", 9, TXT), I("Notes", "Notes", 22),
    A("TotalJobPrice", "Their Total Job Price $", 'IF(OR([Price]="",AND([Basis]="Hourly",N([Qty])=0)),"",[Price]*IF(N([Qty])=0,1,[Qty]))', 11, USD0),
    A("ImpliedHourly", "Their Effective $/Hr", 'IF([Price]="","",IF([Basis]="Hourly",[Price],IFERROR([TotalJobPrice]/IF(N([OurHours])>0,[OurHours],INDEX([Project Types!AvgHours],MATCH([ProjectType],[Project Types!Label],0))),"")))', 10, USD),
    A("PriceGap", "Our Price vs Theirs %", 'IF(OR(N([OurPrice])=0,N([TotalJobPrice])=0),"",[OurPrice]/[TotalJobPrice]-1)', 9, PCT, note="Negative = we were cheaper."),
    A("Recent", "Recent? (helper)", 'IF([Date]>=EDATE(TODAY(),-MarketLookback),1,0)', 8, INT, note="1 = within the look-back window on Settings. Only recent prices feed the analysis."),
]
register("Market Rates", market)

# ================================================================ ANALYSIS COLUMN SPECS
def M(key, header, f, fmt=None, width=12, total=None, note=None):
    return dict(key=key, header=header, f=f, fmt=fmt, width=width, total=total, note=note)

VERDICT = ('IF([Count]=0,"No data yet",IF([Count]<MinSample,"Need more data ("&[Count]&" of "&MinSample&")",'
           'IF([GPHr]="","Log hours to evaluate",IF([GPHr]<OverheadPerHour,"⚠ Below overhead — reprice or drop",'
           'IF(AND([GPHr]>=TargetGPPerHour,[GM]>=TargetGM),"★ Prioritize & grow","OK — tighten time & cost")))))')

def group_cols(pk, qk=None, mk=None, extra_after_label=None, market_price=False):
    P = lambda k: f"[Projects!{k}]"
    base = f'{P(pk)},[Label],{P("InPeriod")},1'
    cols = [M("Label", "", None, None, 26)]
    if extra_after_label:
        cols += extra_after_label
    cols += [
        M("Count", "Completed Projects", f"COUNTIFS({base})", INT, 10, "SUM"),
        M("Rev", "Revenue $", f'SUMIFS({P("Revenue")},{base})', USD0, 12, "SUM"),
        M("Cost", "Direct Cost $", f'SUMIFS({P("DirectCost")},{base})', USD0, 12, "SUM"),
        M("GP", "Gross Profit $", f'SUMIFS({P("GrossProfit")},{base})', USD0, 12, "SUM"),
        M("GM", "Gross Margin %", 'IF([Rev]=0,"",[GP]/[Rev])', PCT, 9, "RATIO:GP/Rev"),
        M("Hours", "Labor Hours", f'SUMIFS({P("LaborHours")},{base})', NUM, 10, "SUM"),
        M("RevHr", "Revenue / Labor Hr", 'IF([Hours]=0,"",[Rev]/[Hours])', USD, 10, "RATIO:Rev/Hours"),
        M("GPHr", "Gross Profit / Labor Hr", 'IF([Hours]=0,"",[GP]/[Hours])', USD, 10, "RATIO:GP/Hours",
          note="The #1 number for choosing work: profit produced per crew hour."),
        M("Rank", "Rank (GP / Hr)", 'IF([GPHr]="","",RANK([GPHr],{rng:GPHr}))', INT, 7),
        M("AvgRev", "Avg Revenue / Project", 'IF([Count]=0,"",[Rev]/[Count])', USD0, 11, "RATIO:Rev/Count"),
        M("AvgGP", "Avg Gross Profit / Project", 'IF([Count]=0,"",[GP]/[Count])', USD0, 11, "RATIO:GP/Count"),
        M("AvgHours", "Avg Labor Hrs / Project", 'IF([Count]=0,"",[Hours]/[Count])', "0.0", 10, "RATIO:Hours/Count"),
        M("AvgDays", "Avg Calendar Days", f'IFERROR(AVERAGEIFS({P("CalendarDays")},{base}),"")', "0.0", 9),
        M("HrsVar", "Hours vs Estimate %", f'IFERROR(SUMIFS({P("LaborHours")},{base},{P("EstHours")},">0")/SUMIFS({P("EstHours")},{base},{P("EstHours")},">0")-1,"")', PCT, 9,
          note="Weighted: total actual hours ÷ total estimated hours − 1. Positive = you under-estimate this work."),
        M("CostVar", "Cost vs Estimate %", f'IFERROR(SUMIFS({P("DirectCost")},{base},{P("EstCost")},">0")/SUMIFS({P("EstCost")},{base},{P("EstCost")},">0")-1,"")', PCT, 9),
        M("Wrench", "Wrench-Time %", f'IFERROR(SUMIFS({P("ProductiveHours")},{base})/[Hours],"")', PCT, 9),
    ]
    if qk:
        qb = f'[Quotes!{qk}],[Label],[Quotes!InPeriod],1'
        cols += [
            M("Quotes", "Quotes", f"COUNTIFS({qb})", INT, 8, "SUM"),
            M("WinRate", "Win Rate %", f'IFERROR(COUNTIFS({qb},[Quotes!Status],"Won")/(COUNTIFS({qb},[Quotes!Status],"Won")+COUNTIFS({qb},[Quotes!Status],"Lost")+COUNTIFS({qb},[Quotes!Status],"Expired")),"")', PCT, 8,
              note="Won ÷ (Won + Lost + Expired). High win rate + high GP/hr = you may be under-priced."),
        ]
    cols += [
        M("CBRate", "Callbacks / Project", f'IF([Count]=0,"",SUMIFS({P("Callbacks")},{base})/[Count])', "0.00", 9),
        M("Sat", "Avg Satisfaction", f'IFERROR(AVERAGEIFS({P("Satisfaction")},{base}),"")', "0.0", 9),
        M("Share", "Share of Total Gross Profit", 'IFERROR([GP]/SUM({rng:GP}),"")', PCT, 9),
    ]
    if mk:
        mb = f'[Market Rates!{mk}],[Label],[Market Rates!Recent],1,[Market Rates!Basis],"<>Hourly"'
        if market_price:
            cols += [
                M("MktPrice", "Market Avg Job Price $", f'IFERROR(AVERAGEIFS([Market Rates!TotalJobPrice],{mb}),"")', USD0, 11),
                M("PriceVsMkt", "Our Avg Price vs Market %", 'IF(OR([AvgRev]="",[MktPrice]=""),"",[AvgRev]/[MktPrice]-1)', PCT, 9),
            ]
        cols += [
            M("MktHr", "Market $/Hr (competitor job prices)", f'IFERROR(AVERAGEIFS([Market Rates!ImpliedHourly],{mb}),"")', USD, 10,
              note="Competitors' whole-job prices ÷ the hours that job takes you. Same basis as your Revenue / Labor Hr (both include materials). Posted hourly rates are compared separately on Pricing & Hiring."),
            M("VsMkt", "Our Revenue/Hr vs Market %", 'IF(OR([RevHr]="",[MktHr]=""),"",[RevHr]/[MktHr]-1)', PCT, 9,
              note="Negative = competitors bring in more per hour for this work than you do."),
        ]
    cols += [M("Verdict", "Verdict", VERDICT, None, 30,
               note="★ = profit per hour beats your target AND margin beats target. ⚠ = doesn't even cover overhead. Thresholds come from Settings and Pricing & Hiring.")]
    return cols

trade_cols = group_cols("Trade", "Trade", "Trade")
register("Trade Analysis", trade_cols)
ptype_cols = group_cols("ProjectType", "ProjectType", "ProjectType",
                        extra_after_label=[M("PTTrade", "Trade", None, None, 13)], market_price=True)
register("Project Types", ptype_cols)

# ================================================================ SETTINGS CONTENT
TRADES = ["HVAC", "Plumbing", "Electrical", "Flooring", "Remodeling", "Refurbishing", "General / Handyman"]
PTYPES = [
    ("AC Replacement", "HVAC"), ("Furnace Replacement", "HVAC"), ("Heat Pump Install", "HVAC"),
    ("Full HVAC System (New Install)", "HVAC"), ("Ductless Mini-Split Install", "HVAC"), ("Ductwork Install / Repair", "HVAC"),
    ("HVAC Repair / Service Call", "HVAC"), ("HVAC Maintenance / Tune-Up", "HVAC"), ("Thermostat Install", "HVAC"),
    ("Water Heater Replacement (Tank)", "Plumbing"), ("Tankless Water Heater Install", "Plumbing"),
    ("Fixture Install / Replace", "Plumbing"), ("Leak Repair", "Plumbing"), ("Repipe", "Plumbing"),
    ("Drain / Sewer Line", "Plumbing"), ("Plumbing Repair / Service Call", "Plumbing"),
    ("Panel Upgrade", "Electrical"), ("Circuit / Outlet Addition", "Electrical"), ("Lighting Install", "Electrical"),
    ("Ceiling Fan Install", "Electrical"), ("EV Charger Install", "Electrical"), ("Rewire", "Electrical"),
    ("Electrical Repair / Troubleshoot", "Electrical"),
    ("LVP / Laminate Install", "Flooring"), ("Hardwood Install", "Flooring"), ("Hardwood Refinishing", "Flooring"),
    ("Tile Floor Install", "Flooring"), ("Carpet Install", "Flooring"), ("Subfloor Repair", "Flooring"),
    ("Kitchen Remodel", "Remodeling"), ("Bathroom Remodel", "Remodeling"), ("Basement Finish", "Remodeling"),
    ("Addition", "Remodeling"), ("Whole-Home Remodel", "Remodeling"),
    ("Rental Turnover / Make-Ready", "Refurbishing"), ("Pre-Sale Refresh", "Refurbishing"),
    ("Interior Painting", "Refurbishing"), ("Exterior Painting", "Refurbishing"), ("Drywall Repair", "Refurbishing"),
    ("General Handyman", "General / Handyman"), ("Door / Window Install", "General / Handyman"),
    ("Deck / Fence Repair", "General / Handyman"), ("Carpentry / Trim", "General / Handyman"),
]
TASKS = [("Site Visit / Estimate", "No"), ("Material Pickup", "No"), ("Travel", "No"),
         ("Demo / Tear-Out", "Yes"), ("Prep / Protection", "Yes"), ("Rough-In", "Yes"), ("Install", "Yes"),
         ("Repair / Service", "Yes"), ("Finish / Trim-Out", "Yes"), ("Testing / Inspection", "Yes"),
         ("Cleanup / Haul-Away", "Yes"), ("Warranty / Callback", "No"), ("Waiting / Delay", "No"),
         ("Admin / Office", "No"), ("Training", "No"), ("Shop / Vehicle Maintenance", "No")]
JOB_CATS = ["Materials", "Subcontractor", "Equipment Rental", "Permits & Inspections", "Disposal / Dumpster",
            "Job Supplies & Consumables", "Job Fuel / Delivery", "Other Job Cost"]
OH_CATS = ["Insurance", "Vehicle Payment & Maintenance", "Fuel (General)", "Tools & Equipment", "Phone & Internet",
           "Software & Subscriptions", "Marketing & Advertising", "Office & Admin", "Accounting & Legal",
           "Licenses & Certifications", "Training & Education", "Rent / Storage / Shop", "Bank & Card Fees",
           "Office Payroll (non-field)", "Other Overhead"]
LEAD_SOURCES = ["Repeat Client", "Referral – Client", "Referral – Realtor / Property Mgr", "Google Search / Maps",
                "Google Local Services Ads", "Website", "Facebook / Instagram", "Nextdoor", "Yelp",
                "Angi / HomeAdvisor", "Thumbtack", "Yard Sign / Truck Wrap", "Flyer / Door Hanger",
                "Walk-in / Phone", "Other"]
if GSHEETS:
    LEAD_SOURCES.insert(14, "WhatsApp")
CLIENT_TYPES = ["Homeowner", "Landlord / Investor", "Property Manager", "Realtor", "Commercial / Business",
                "General Contractor", "HOA"]

LISTS = [
    ("Trades", "Trades", TRADES),
    ("ProjectTypes", "Project Types → Trade", PTYPES),
    ("Tasks", "Tasks → Wrench Time?", TASKS),
    ("ExpCategories", "Expense Categories → Type", [(c, "Job Cost") for c in JOB_CATS] + [(c, "Overhead") for c in OH_CATS]),
    ("LeadSources", "Lead Sources", LEAD_SOURCES),
    ("ClientTypes", "Client Types", CLIENT_TYPES),
    ("PropertyTypes", "Property Types", ["Single-Family", "Townhouse", "Condo / Apartment", "Multi-Family (2-4)",
                                         "Multi-Family (5+)", "Mobile / Manufactured", "Commercial – Retail",
                                         "Commercial – Office", "Commercial – Restaurant", "Other"]),
    ("Occupancy", "Occupancy", ["Owner-Occupied", "Tenant-Occupied", "Vacant"]),
    ("ProjectStatuses", "Project Statuses", ["Scheduled", "In Progress", "On Hold", "Completed", "Cancelled"]),
    ("QuoteStatuses", "Quote Statuses", ["Draft", "Sent", "Follow-Up", "Won", "Lost", "Expired", "Declined by Us"]),
    ("ContractTypes", "Contract Types", ["Fixed Price", "Time & Materials", "Cost Plus", "Service Call (Flat Fee)"]),
    ("LostReasons", "Lost Reasons", ["Price Too High", "Went With Competitor", "Timing / Availability", "No Response",
                                     "Project Postponed", "Scope Changed", "Financing Fell Through", "Doing It Themselves",
                                     "Not a Good Fit", "Other"]),
    ("COReasons", "Change Order Reasons", ["Client Request", "Hidden / Unforeseen Condition", "Code Requirement",
                                           "Design Change", "Material Change", "Our Error (No Charge)"]),
    ("COStatuses", "CO Statuses", ["Pending", "Approved", "Rejected"]),
    ("PaymentMethods", "Payment Methods", ["Check", "Cash", "Zelle", "ACH / Bank Transfer", "Credit / Debit Card",
                                           "Financing", "Venmo / PayPal", "Other"]),
    ("InvoiceTypes", "Invoice Types", ["Deposit", "Progress", "Final", "Change Order", "Service Call"]),
    ("PaidWith", "Paid With", ["Business Card", "Business Checking", "Cash", "Personal (Reimburse)", "Vendor Account / Net Terms"]),
    ("Roles", "Roles", ["Owner", "Lead Technician", "Technician", "Apprentice", "Helper / Laborer", "Office / Admin"]),
    ("EmpTypes", "Employment Types", ["Owner", "W-2 Employee", "1099 Contractor"]),
    ("Skill", "Skill Levels", ["Helper", "Apprentice", "Journeyman", "Master", "Specialist"]),
    ("WorkerStatus", "Worker Status", ["Active", "Inactive"]),
    ("VendorTypes", "Vendor Types", ["Supplier", "Subcontractor", "Equipment Rental", "Disposal", "Service Provider", "Other"]),
    ("CallbackCauses", "Callback Root Causes", ["Workmanship", "Installation Oversight", "Material / Product Defect",
                                                "Client Misuse", "Pre-Existing Condition", "Design Issue", "Other"]),
    ("PricingBasis", "Pricing Basis", ["Hourly", "Flat Price (Whole Job)", "Per Sq Ft", "Per Unit / Fixture"]),
    ("MarketSources", "Market Data Sources", ["Client Showed Competitor Quote", "Lost-Quote Feedback", "Competitor Website",
                                              "Phone / Mystery Shop", "Angi / HomeAdvisor / Thumbtack",
                                              "Industry Cost Guide", "Supplier / Industry Contact", "Other"]),
    ("Urgency", "Urgency", ["Emergency (same/next day)", "Soon (within 2 weeks)", "Flexible / Planning"]),
    ("Contact", "Contact Methods", ["Call", "Text", "Email", "WhatsApp"]),
    ("Foundation", "Foundation", ["Slab", "Crawlspace", "Basement (Unfinished)", "Basement (Finished)", "Pier & Beam"]),
    ("HVACTypes", "HVAC Types", ["Central AC + Gas Furnace", "Heat Pump", "Ductless Mini-Split", "Boiler / Radiant",
                                "Electric Baseboard", "Window / PTAC", "None / Unknown"]),
    ("WHTypes", "Water Heater Types", ["Tank – Gas", "Tank – Electric", "Tankless – Gas", "Tankless – Electric",
                                       "Heat Pump Water Heater", "Unknown"]),
    ("Pipe", "Pipe Materials", ["Copper", "PEX", "CPVC", "Galvanized Steel", "Polybutylene", "Mixed / Unknown"]),
    ("Sensitivity", "Price Sensitivity", ["Low", "Medium", "High"]),
    ("YesNo", "Yes / No", ["Yes", "No"]),
    ("YesNoNA", "Yes / No / N/A", ["Yes", "No", "N/A"]),
    ("Zones", "Travel Zones (auto)", ["=\"Zone 1 (0–\"&Zone1Max&\" mi)\"", "=\"Zone 2 (\"&Zone1Max&\"–\"&Zone2Max&\" mi)\"",
                                      "=\"Zone 3 (\"&Zone2Max&\"–\"&Zone3Max&\" mi)\"", "=\"Zone 4 (\"&Zone3Max&\"+ mi)\""]),
]
LIST_ROWS = 150

ASSUMPTIONS = [
    ("BizName", "Business name", "Your Company Name", None, "Shown on the Dashboard."),
    ("BurdenDefault", "Default labor burden %", 0.25, PCT,
     "Extra cost on top of hourly pay: employer payroll taxes (~7.65% FICA + ~1–3% unemployment), workers' comp (varies a lot by trade, often 5–15%+), benefits. 25% is a starting estimate — replace with your payroll provider / insurance agent's real numbers."),
    ("TargetGM", "Target gross margin %", 0.40, PCT,
     "Revenue minus direct job costs, as % of revenue. Residential trade & remodel businesses commonly target ~35–50%. Your goal — adjust."),
    ("TargetNet", "Target net profit %", 0.15, PCT,
     "Profit left after overhead. 10–20% is a common goal for a healthy contractor. Your goal — adjust."),
    ("DefaultTerms", "Default invoice terms (days)", 7, INT, "Used when an invoice has no terms entered."),
    ("MinSample", "Min. completed projects before a verdict", 3, INT,
     "Analysis tabs say 'Need more data' below this. Don't make big decisions on 1–2 jobs."),
    ("MarketLookback", "Market price look-back (months)", 24, INT, "Competitor prices older than this are ignored."),
    ("HVACAgeFlag", "Flag HVAC equipment older than (yrs)", 12, INT, "Typical AC/furnace life is roughly 12–20 years."),
    ("WHAgeFlag", "Flag water heater older than (yrs)", 10, INT, "Typical tank water heater life is roughly 8–12 years."),
    ("PanelFlagAmps", "Flag electrical panels at or below (amps)", 100, INT, "100A and smaller panels are common upgrade candidates (EVs, heat pumps, additions)."),
    ("TierA", "Client Tier A — lifetime gross profit ≥ $", 10000, USD0, "Your best clients."),
    ("TierB", "Client Tier B — lifetime gross profit ≥ $", 3000, USD0, "Everyone below this is Tier C."),
    ("Zone1Max", "Travel Zone 1 ends at (miles)", 10, INT, "One-way miles from your shop / home base."),
    ("Zone2Max", "Travel Zone 2 ends at (miles)", 25, INT, ""),
    ("Zone3Max", "Travel Zone 3 ends at (miles)", 40, INT, "Anything farther is Zone 4."),
]

# ================================================================ BUILD
wb = Workbook()
wb.remove(wb.active)
ORDER = ["Start Here", "Dashboard", "Clients", "Properties", "Quotes", "Projects", "Change Orders", "Time Log",
         "Expenses", "Invoices", "Team", "Vendors", "Callbacks", "Market Rates", "Trade Analysis", "Project Types",
         "Time Analysis", "Pricing & Hiring", "Marketing", "Clients & Areas", "Monthly Trend", "Settings"]
if GSHEETS:
    ORDER.append("WhatsApp Log")
WS = {n: wb.create_sheet(n) for n in ORDER}
TAB = {"Start Here": "000000", "Dashboard": GOLD, "Settings": "808080", "WhatsApp Log": "25D366"}
for n in ["Clients", "Properties", "Quotes", "Projects", "Change Orders", "Time Log", "Expenses", "Invoices",
          "Team", "Vendors", "Callbacks", "Market Rates"]:
    TAB[n] = "2F5597"
for n in ["Trade Analysis", "Project Types", "Time Analysis", "Pricing & Hiring", "Marketing", "Clients & Areas", "Monthly Trend"]:
    TAB[n] = "548235"
for n, c in TAB.items():
    if n in WS:
        WS[n].sheet_properties.tabColor = c
for ws in WS.values():
    ws.sheet_view.showGridLines = False if ws.title in ("Start Here", "Dashboard", "Pricing & Hiring") else True
    ws.sheet_view.zoomScale = 100

NAMES = {}

def name(n, ref):
    NAMES[n] = ref
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


def title(ws, t, sub):
    ws["A1"] = t
    ws["A1"].font = f_title
    ws["A2"] = sub
    ws["A2"].font = f_sub
    ws.row_dimensions[1].height = 24


def note(cell, text, w=320, h=110):
    c = Comment(text, "Workbook guide")
    c.width, c.height = w, h
    cell.comment = c


# ---------------- Settings
ws = WS["Settings"]
title(ws, "Settings & Lists", "Your business assumptions (yellow = set these) and every dropdown list in the workbook. Add items to the bottom of any list and they appear in the dropdowns.")
ws["A4"], ws["B4"], ws["C4"] = "Assumption", "Value", "What it means / where the default comes from"
for c in ("A4", "B4", "C4"):
    ws[c].font, ws[c].fill, ws[c].alignment = f_hdr, fill_in_h, wrap_c
for i, (n, label, val, fmt, nt) in enumerate(ASSUMPTIONS):
    r = FIRST + i
    ws.cell(r, 1, label).font = f_bold
    c = ws.cell(r, 2, val)
    c.font, c.fill, c.border = f_in, fill_yellow, box
    if fmt:
        c.number_format = fmt
    t = ws.cell(r, 3, nt)
    t.font, t.alignment = f_body, wrap_l
    ws.row_dimensions[r].height = 42 if len(nt) > 90 else (28 if len(nt) > 45 else 15)
    name(n, f"Settings!$B${r}")
ws.column_dimensions["A"].width = 38
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 60
ws.column_dimensions["D"].width = 3

col = 5
LIST_POS = {}
for lname, header, vals in LISTS:
    two = isinstance(vals[0], tuple)
    L1 = GL(col)
    ws.cell(HDR_ROW, col, header)
    span = 2 if two else 1
    for k in range(span):
        h = ws.cell(HDR_ROW, col + k)
        h.font, h.fill, h.alignment = f_hdr, fill_auto_h if lname == "Zones" else fill_in_h, wrap_c
    if two:
        ws.merge_cells(start_row=HDR_ROW, start_column=col, end_row=HDR_ROW, end_column=col + 1)
    for i, v in enumerate(vals):
        if two:
            ws.cell(FIRST + i, col, v[0]).font = f_in
            ws.cell(FIRST + i, col + 1, v[1]).font = f_in
        else:
            c = ws.cell(FIRST + i, col, v)
            c.font = f_auto if lname == "Zones" else f_in
    if lname == "Zones":
        name("List_Zones", f"Settings!${L1}${FIRST}:${L1}${FIRST + 3}")
    elif GSHEETS:
        name(f"List_{lname}", f"Settings!${L1}${FIRST}:${L1}${FIRST + LIST_ROWS - 1}")
    else:
        name(f"List_{lname}", f"OFFSET(Settings!${L1}${FIRST},0,0,MAX(1,COUNTA(Settings!${L1}${FIRST}:${L1}${FIRST + LIST_ROWS - 1})),1)")
    LIST_POS[lname] = (col, len(vals))
    ws.column_dimensions[L1].width = 30 if lname in ("ProjectTypes", "ExpCategories", "LeadSources", "MarketSources") else 22
    if two:
        L2 = GL(col + 1)
        ws.column_dimensions[L2].width = 12
        rng1 = f"Settings!${L1}${FIRST}:${L1}${FIRST + LIST_ROWS - 1}"
        rng2 = f"Settings!${L2}${FIRST}:${L2}${FIRST + LIST_ROWS - 1}"
        pair = {"ProjectTypes": ("PT_Types", "PT_Trades"), "Tasks": ("Task_Names", "Task_Productive"),
                "ExpCategories": ("Cat_Names", "Cat_Types")}[lname]
        name(pair[0], rng1)
        name(pair[1], rng2)
    col += span
ws.freeze_panes = "A5"
ws.row_dimensions[HDR_ROW].height = 32
pt_col = LIST_POS["ProjectTypes"][0]
dv = DataValidation(type="list", formula1="=List_Trades", allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"{GL(pt_col + 1)}{FIRST}:{GL(pt_col + 1)}{FIRST + LIST_ROWS - 1}")
dv2 = DataValidation(type="list", formula1='"Job Cost,Overhead"', allow_blank=True)
ws.add_data_validation(dv2)
ec = LIST_POS["ExpCategories"][0]
dv2.add(f"{GL(ec + 1)}{FIRST}:{GL(ec + 1)}{FIRST + LIST_ROWS - 1}")
dv3 = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
ws.add_data_validation(dv3)
tc = LIST_POS["Tasks"][0]
dv3.add(f"{GL(tc + 1)}{FIRST}:{GL(tc + 1)}{FIRST + LIST_ROWS - 1}")


def settings_cell(lname, i, second=False):
    c0 = LIST_POS[lname][0] + (1 if second else 0)
    return f"Settings!${GL(c0)}${FIRST + i}"


# ---------------- Data sheets
DATA = {
    "Clients": dict(cols=clients, rows=500, prefix="C-", digits=4,
                    sub="One row per client (person or company). Everything about who you work for."),
    "Properties": dict(cols=properties, rows=500, prefix="P-", digits=4,
                       sub="One row per job site. What you learn on site today is tomorrow's sales list (see Future Work Flags)."),
    "Quotes": dict(cols=quotes, rows=1000, prefix="Q-", digits=4,
                   sub="Every estimate you give — won, lost or pending. This is how you learn your win rate and whether you're priced right."),
    "Projects": dict(cols=projects, rows=500, prefix="J-", digits=4,
                     sub="One row per job. You fill the blue columns; everything gray calculates from Time Log, Expenses, Invoices, Change Orders and Callbacks."),
    "Change Orders": dict(cols=change_orders, rows=300, prefix="CO-", digits=4,
                          sub="Any change to scope or price after the job was agreed. Only Approved change orders count toward revenue."),
    "Time Log": dict(cols=time_log, rows=3000, prefix=None,
                     sub="Every paid hour, by person, job and task. The most important tab — without it you can't know what your time is worth."),
    "Expenses": dict(cols=expenses, rows=3000, prefix=None,
                     sub="Every dollar out. Tag it to a Project ID (job cost) or OVERHEAD (general business cost)."),
    "Invoices": dict(cols=invoices, rows=1000, prefix="INV-", digits=4,
                     sub="Every invoice and payment. Tracks cash in, who owes you, and how long clients take to pay."),
    "Team": dict(cols=team, rows=30, prefix="W-", digits=2,
                 sub="Everyone who works jobs — including owners. Pay rates here drive every labor cost in the workbook."),
    "Vendors": dict(cols=vendors, rows=150, prefix="V-", digits=3,
                    sub="Suppliers and subcontractors: terms, discounts, insurance, and how much you spend with each."),
    "Callbacks": dict(cols=callbacks, rows=200, prefix="CB-", digits=3,
                      sub="Any return trip to fix finished work. Every callback is profit lost — track root causes to stop repeats."),
    "Market Rates": dict(cols=market, rows=500, prefix=None,
                         sub="Competitor prices you learn about. Feeds the market comparison on Trade Analysis, Project Types and Pricing & Hiring."),
}

EXAMPLES = {
    "Clients": [dict(ClientID="C-0001", ClientType="Homeowner", FirstName="Sarah", LastName="Example", Phone="555-010-0142",
                     Email="sarah@example.com", PrefContact="Text", Language="English", Address="123 Maple St",
                     City="Springfield", State="IL", ZIP="62701", LeadSource="Google Search / Maps",
                     FirstContact=D(2026, 3, 2), OptIn="Yes", PriceSens="Medium", DNS="No",
                     Notes="EXAMPLE ROW — delete once you add real data")],
    "Properties": [dict(PropertyID="P-0001", ClientID="C-0001", Street="123 Maple St", City="Springfield", State="IL",
                        ZIP="62701", County="Sangamon", Area="West Side", DistanceMi=12, DriveMin=22,
                        PropertyType="Single-Family", Occupancy="Owner-Occupied", YearBuilt=1994, SqFt=1850, Stories=2,
                        Beds=3, Baths=2.5, Foundation="Basement (Unfinished)", HVACType="Central AC + Gas Furnace",
                        HVACYear=2010, WHType="Tank – Gas", WHYear=2014, PanelAmps=150, PanelBrand="Square D",
                        Pipe="PEX", HOA="No", Jurisdiction="City of Springfield", Access="Park in driveway",
                        Pets="1 dog", Notes="EXAMPLE ROW — delete once you add real data")],
    "Quotes": [dict(QuoteID="Q-0001", DateReceived=D(2026, 3, 2), ClientID="C-0001", PropertyID="P-0001",
                    LeadSource="Google Search / Maps", ProjectType="AC Replacement",
                    Scope="Replace 3-ton AC condenser + evaporator coil", Urgency="Soon (within 2 weeks)",
                    SiteVisit=D(2026, 3, 4), QuoteSent=D(2026, 3, 5), EstimatingHrs=2, EstLaborHrs=24,
                    EstMaterials=3800, EstSubs=0, EstOther=250, QuotedPrice=7900, CompetitorName="Competitor A",
                    CompetitorPrice=8600, Status="Won", DecisionDate=D(2026, 3, 9), ProjectID="J-0001", FollowUps=1,
                    Notes="EXAMPLE ROW — delete once you add real data")],
    "Projects": [dict(ProjectID="J-0001", QuoteID="Q-0001", ClientID="C-0001", PropertyID="P-0001",
                      ProjectName="Example AC replacement", ProjectType="AC Replacement", Status="Completed",
                      ContractType="Fixed Price", Complexity=3, Permit="Yes", LeadTech="W-02",
                      SchedStart=D(2026, 3, 16), SchedEnd=D(2026, 3, 17), ActualStart=D(2026, 3, 16),
                      ActualEnd=D(2026, 3, 18), ContractValue=7900, Inspection="Yes", Satisfaction=9,
                      Testimonial="Yes", Review="Yes", Notes="EXAMPLE ROW — delete once you add real data")],
    "Change Orders": [dict(COID="CO-0001", ProjectID="J-0001", Date=D(2026, 3, 17),
                           Description="Replace corroded refrigerant line set", Reason="Hidden / Unforeseen Condition",
                           AddedHours=3, AddedCost=180, Price=650, Status="Approved", ApprovedDate=D(2026, 3, 17),
                           Signed="Yes", Notes="EXAMPLE ROW")],
    "Time Log": [
        dict(Date=D(2026, 3, 4), WorkerID="W-01", ProjectID="OVERHEAD", Task="Site Visit / Estimate", ManualHours=2, Notes="EXAMPLE — estimate for Q-0001"),
        dict(Date=D(2026, 3, 16), WorkerID="W-02", ProjectID="J-0001", Task="Material Pickup", Start=T(7, 0), End=T(8, 0), Notes="EXAMPLE"),
        dict(Date=D(2026, 3, 16), WorkerID="W-02", ProjectID="J-0001", Task="Travel", Start=T(8, 0), End=T(8, 30), Notes="EXAMPLE"),
        dict(Date=D(2026, 3, 16), WorkerID="W-01", ProjectID="J-0001", Task="Travel", Start=T(8, 0), End=T(8, 30), Notes="EXAMPLE"),
        dict(Date=D(2026, 3, 16), WorkerID="W-02", ProjectID="J-0001", Task="Demo / Tear-Out", Start=T(8, 30), End=T(12, 0), Notes="EXAMPLE"),
        dict(Date=D(2026, 3, 16), WorkerID="W-01", ProjectID="J-0001", Task="Demo / Tear-Out", Start=T(8, 30), End=T(12, 0), Notes="EXAMPLE"),
        dict(Date=D(2026, 3, 16), WorkerID="W-02", ProjectID="J-0001", Task="Install", Start=T(12, 30), End=T(17, 0), Notes="EXAMPLE"),
        dict(Date=D(2026, 3, 17), WorkerID="W-02", ProjectID="J-0001", Task="Install", Start=T(8, 0), End=T(16, 30), BreakMin=30, Notes="EXAMPLE"),
        dict(Date=D(2026, 3, 17), WorkerID="W-01", ProjectID="J-0001", Task="Install", Start=T(8, 0), End=T(16, 30), BreakMin=30, Notes="EXAMPLE"),
        dict(Date=D(2026, 3, 18), WorkerID="W-02", ProjectID="J-0001", Task="Testing / Inspection", Start=T(9, 0), End=T(11, 0), Notes="EXAMPLE"),
    ],
    "Expenses": [
        dict(Date=D(2026, 3, 1), ProjectID="OVERHEAD", Category="Insurance", Vendor="Insurance Co.", Description="General liability — monthly", Amount=310, Tax=0, PaidWith="Business Checking", Notes="EXAMPLE"),
        dict(Date=D(2026, 3, 5), ProjectID="OVERHEAD", Category="Marketing & Advertising", Vendor="Google Ads", Description="Search ads — March", Amount=250, Tax=0, PaidWith="Business Card", Channel="Google Search / Maps", Notes="EXAMPLE"),
        dict(Date=D(2026, 3, 16), ProjectID="J-0001", Category="Materials", Vendor="Metro HVAC Supply", Description="3-ton condenser + coil", Amount=3650, Tax=292, PaidWith="Vendor Account / Net Terms", Notes="EXAMPLE"),
        dict(Date=D(2026, 3, 16), ProjectID="J-0001", Category="Permits & Inspections", Vendor="City of Springfield", Description="Mechanical permit", Amount=175, Tax=0, PaidWith="Business Card", Notes="EXAMPLE"),
        dict(Date=D(2026, 3, 17), ProjectID="J-0001", Category="Materials", Vendor="Metro HVAC Supply", Description="Line set (change order)", Amount=160, Tax=12.8, PaidWith="Vendor Account / Net Terms", Notes="EXAMPLE"),
        dict(Date=D(2026, 3, 18), ProjectID="J-0001", Category="Disposal / Dumpster", Vendor="Metro HVAC Supply", Description="Old unit disposal / refrigerant recovery", Amount=60, Tax=0, PaidWith="Business Card", Notes="EXAMPLE"),
    ],
    "Invoices": [
        dict(InvoiceNo="INV-0001", ProjectID="J-0001", InvoiceType="Deposit", InvoiceDate=D(2026, 3, 9), Terms=0, Amount=3950, AmountPaid=3950, DatePaid=D(2026, 3, 10), Method="Zelle", Notes="EXAMPLE"),
        dict(InvoiceNo="INV-0002", ProjectID="J-0001", InvoiceType="Final", InvoiceDate=D(2026, 3, 18), Terms=7, Amount=4600, AmountPaid=4600, DatePaid=D(2026, 3, 24), Method="Credit / Debit Card", Fee=133.4, Notes="EXAMPLE"),
    ],
    "Team": [
        dict(WorkerID="W-01", Name="Owner — edit name", Role="Owner", EmpType="Owner", Skill="Journeyman", Status="Active", BasePay=30,
             Notes="Base pay = what you'd pay someone to replace you on the tools. PLACEHOLDER — set your own."),
        dict(WorkerID="W-02", Name="Partner — edit name", Role="Owner", EmpType="Owner", Skill="Master", Status="Active", BasePay=45,
             Notes="PLACEHOLDER rate — set your own."),
    ],
    "Vendors": [
        dict(VendorID="V-001", VendorName="Metro HVAC Supply", VendorType="Supplier", Specialty="HVAC", Terms="Net 30", Discount=0.1, W9="Yes", Notes="EXAMPLE ROW"),
        dict(VendorID="V-002", VendorName="Google Ads", VendorType="Service Provider", Notes="EXAMPLE ROW"),
    ],
    "Callbacks": [],
    "Market Rates": [
        dict(Date=D(2026, 3, 9), Competitor="Competitor A", Trade="HVAC", ProjectType="AC Replacement", Basis="Flat Price (Whole Job)",
             Price=8600, OurHours=24, OurPrice=7900, Source="Client Showed Competitor Quote", AreaZip="62701", Notes="EXAMPLE ROW"),
        dict(Date=D(2026, 4, 2), Competitor="Competitor B", Trade="Electrical", Basis="Hourly", Price=95,
             Source="Phone / Mystery Shop", AreaZip="62701", Notes="EXAMPLE ROW — their posted hourly rate"),
    ],
}

YELLOW_EX = {("Team", "BasePay")}

for sname, spec in DATA.items():
    ws = WS[sname]
    cols, nrows = spec["cols"], spec["rows"]
    title(ws, sname, spec["sub"])
    last = FIRST + nrows - 1
    if spec.get("prefix"):
        ws["A3"] = f'="Next ID: {spec["prefix"]}"&TEXT(COUNTA($A${FIRST}:$A${last})+1,"{"0" * spec["digits"]}")'
        ws["A3"].font = Font(name=FONT, size=10, bold=True, color=GREEN_T)
    ws["D3"] = "Blue header = you type it     Gray header = calculates automatically (don't type there)     Hover a header's red corner for tips"
    ws["D3"].font = f_sub
    guard = cols[0]["key"]
    lt = SHEETS[sname]["letters"]
    exrows = EXAMPLES.get(sname, [])
    for j, c in enumerate(cols):
        L = GL(j + 1)
        h = ws.cell(HDR_ROW, j + 1, c["header"])
        h.font, h.alignment, h.border = f_hdr, wrap_c, box
        h.fill = fill_in_h if c["kind"] == "in" else fill_auto_h
        if c.get("note"):
            note(h, c["note"])
        ws.column_dimensions[L].width = c["width"]
        if c["kind"] == "in" and c.get("dv"):
            d = c["dv"]
            if isinstance(d, tuple) and d[0] == "whole":
                v = DataValidation(type="whole", operator="between", formula1=str(d[1]), formula2=str(d[2]), allow_blank=True)
                v.error, v.errorTitle = f"Enter a whole number {d[1]}–{d[2]}", "Invalid value"
            elif isinstance(d, tuple) and d[0] == "free":
                v = DataValidation(type="list", formula1=f"={d[1]}", allow_blank=True, showErrorMessage=False)
            else:
                v = DataValidation(type="list", formula1=f"={d}", allow_blank=True)
                v.error, v.errorTitle = "Pick from the list (add new options on the Settings tab).", "Not in list"
            ws.add_data_validation(v)
            v.add(f"{L}{FIRST}:{L}{last}")
        for r in range(FIRST, last + 1):
            cell = ws.cell(r, j + 1)
            idx = r - FIRST
            if c["kind"] == "in":
                if idx < len(exrows) and c["key"] in exrows[idx]:
                    cell.value = exrows[idx][c["key"]]
                cell.font = f_in
                if (sname, c["key"]) in YELLOW_EX and idx < len(exrows):
                    cell.fill = fill_yellow
            else:
                body = render(c["f"], sname, r)
                cell.value = f'=IF(${lt[guard]}{r}="","",{body})' if c["guard"] else f"={body}"
                cell.font = f_warn if c.get("warn") else f_auto
                cell.fill = fill_auto
            if c.get("fmt"):
                cell.number_format = c["fmt"]
    ws.freeze_panes = ws.cell(FIRST, 2)
    ws.row_dimensions[HDR_ROW].height = 54
    ws.auto_filter.ref = f"A{HDR_ROW}:{GL(len(cols))}{last}"


def idlist(n, sheet, key):
    L = SHEETS[sheet]["letters"][key]
    q = f"'{sheet}'" if " " in sheet else sheet
    if GSHEETS:
        name(n, f"{q}!${L}${FIRST}:${L}${FIRST + DATA[sheet]['rows'] - 1}")
    else:
        name(n, f"OFFSET({q}!${L}${FIRST},0,0,MAX(1,COUNTA({q}!${L}${FIRST}:${L}$5000)),1)")

idlist("List_ClientIDs", "Clients", "ClientID")
idlist("List_PropertyIDs", "Properties", "PropertyID")
idlist("List_ProjectIDs", "Projects", "ProjectID")
idlist("List_WorkerIDs", "Team", "WorkerID")
idlist("List_Vendors", "Vendors", "VendorName")

def cf(sname, key, rules, nrows):
    L = SHEETS[sname]["letters"][key]
    rng = f"{L}{FIRST}:{L}{FIRST + nrows - 1}"
    for formula, fill in rules:
        WS[sname].conditional_formatting.add(rng, FormulaRule(formula=[formula.replace("#", f"{L}{FIRST}")], fill=fill, stopIfTrue=True))

cf("Projects", "GM", [("AND(ISNUMBER(#),#<0)", RED_F), ("AND(ISNUMBER(#),#<TargetGM)", AMB_F), ("ISNUMBER(#)", GRN_F)], 500)
cf("Projects", "HoursVar", [("AND(ISNUMBER(#),#>0.15)", RED_F), ("AND(ISNUMBER(#),#<=0)", GRN_F)], 500)
cf("Projects", "CostVar", [("AND(ISNUMBER(#),#>0.1)", RED_F), ("AND(ISNUMBER(#),#<=0)", GRN_F)], 500)
cf("Projects", "DaysLate", [("AND(ISNUMBER(#),#>0)", RED_F)], 500)
cf("Projects", "GPPerHr", [("AND(ISNUMBER(#),#<OverheadPerHour)", RED_F), ("AND(ISNUMBER(#),#>=TargetGPPerHour)", GRN_F)], 500)
cf("Projects", "Balance", [('AND(ISNUMBER(#),#>0.004)', AMB_F)], 500)
cf("Quotes", "PriceCheck", [('LEFT(#,1)="⚠"', AMB_F), ('LEFT(#,1)="✔"', GRN_F)], 1000)
cf("Quotes", "DaysWaiting", [("AND(ISNUMBER(#),#>=7)", AMB_F)], 1000)
cf("Invoices", "Status", [('#="OVERDUE"', RED_F), ('#="Paid"', GRN_F), ('#="Partial"', AMB_F)], 1000)
cf("Team", "LicenseFlag", [('#="EXPIRED"', RED_F), ('#="Renew soon"', AMB_F)], 30)
cf("Vendors", "COIFlag", [('#="EXPIRED"', RED_F), ('#="Renew soon"', AMB_F)], 150)
cf("Properties", "Opportunity", [('#<>""', GRN_F)], 500)
cf("Clients", "Tier", [('#="A"', GRN_F)], 500)

# ================================================================ ANALYSIS SHEET HELPERS
def hdr(ws, r, c, text, fill=fill_in_h, n=None):
    h = ws.cell(r, c, text)
    h.font, h.fill, h.alignment, h.border = f_hdr, fill, wrap_c, box
    if n:
        note(h, n)
    return h

def section(ws, r, text, c1=1, c2=8):
    ws.cell(r, c1, text).font = f_sec
    for c in range(c1, c2 + 1):
        ws.cell(r, c).fill = fill_sec
    ws.row_dimensions[r].height = 20

def out(ws, r, c, v, fmt=None, bold=False, fill=None, font=None):
    x = ws.cell(r, c, v)
    x.font = font or (f_bold if bold else f_auto)
    x.border = box
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = fill
    return x

def verdict_cf(ws, rng, first_cell):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first_cell},1)="★"'], fill=GRN_F))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first_cell},1)="⚠"'], fill=RED_F))

def build_group(ws_name, cols, label_src, nrows, label_header, sub, extra_label_src=None):
    ws = WS[ws_name]
    title(ws, ws_name, sub)
    ws["A3"] = '="Analysis period: "&AnalysisPeriod&"   (change it on the Dashboard)   ·   Only COMPLETED projects are counted"'
    ws["A3"].font = Font(name=FONT, size=10, bold=True, color=GREEN_T)
    lt = SHEETS[ws_name]["letters"]
    first, last = FIRST, FIRST + nrows - 1
    for j, c in enumerate(cols):
        hdr(ws, HDR_ROW, j + 1, c["header"] or label_header, n=c.get("note"))
        ws.column_dimensions[GL(j + 1)].width = c["width"]
    for i in range(nrows):
        r = FIRST + i
        for j, c in enumerate(cols):
            if c["key"] == "Label":
                out(ws, r, j + 1, f'=IF({label_src(i)}="","",{label_src(i)})', bold=True)
            elif c["key"] == "PTTrade":
                out(ws, r, j + 1, f'=IF($A{r}="","",{extra_label_src(i)})')
            else:
                body = render(c["f"], ws_name, r, first, last)
                out(ws, r, j + 1, f'=IF($A{r}="","",{body})', c["fmt"])
    tr = last + 1
    out(ws, tr, 1, "TOTAL / COMPANY", bold=True, fill=fill_total)
    for j, c in enumerate(cols):
        if j == 0:
            continue
        L = GL(j + 1)
        t = c.get("total")
        v = None
        if t == "SUM":
            v = f"=SUM({L}{first}:{L}{last})"
        elif t and t.startswith("RATIO:"):
            a, b = t[6:].split("/")
            v = f'=IFERROR({lt[a]}{tr}/{lt[b]}{tr},"")'
        out(ws, tr, j + 1, v, c["fmt"], bold=True, fill=fill_total)
    ws.freeze_panes = ws.cell(FIRST, 2)
    ws.row_dimensions[HDR_ROW].height = 48
    L = lt["Verdict"]
    verdict_cf(ws, f"{L}{first}:{L}{last}", f"{L}{first}")
    Lg = lt["GM"]
    ws.conditional_formatting.add(f"{Lg}{first}:{Lg}{last}", FormulaRule(formula=[f"AND(ISNUMBER({Lg}{first}),{Lg}{first}<TargetGM)"], fill=AMB_F))
    Lh = lt["GPHr"]
    ws.conditional_formatting.add(f"{Lh}{first}:{Lh}{last}", FormulaRule(formula=[f"AND(ISNUMBER({Lh}{first}),{Lh}{first}>=TargetGPPerHour)"], fill=GRN_F))
    ws.conditional_formatting.add(f"{Lh}{first}:{Lh}{last}", FormulaRule(formula=[f"AND(ISNUMBER({Lh}{first}),{Lh}{first}<OverheadPerHour)"], fill=RED_F))
    return tr

N_TRADES_ROWS = 12
build_group("Trade Analysis", trade_cols, lambda i: settings_cell("Trades", i), N_TRADES_ROWS, "Trade",
            "Which trades make you the most money per hour of your time — and which don't. Green GP/Hr beats your target; red doesn't cover overhead.")
N_PT_ROWS = 80
build_group("Project Types", ptype_cols, lambda i: settings_cell("ProjectTypes", i), N_PT_ROWS, "Project Type",
            "The detailed version: every kind of job you do, ranked by profit per hour. This is where you decide what to chase and what to reprice.",
            extra_label_src=lambda i: settings_cell("ProjectTypes", i, second=True))

# ---------------- Time Analysis
ws = WS["Time Analysis"]
title(ws, "Time Analysis", "Where the hours actually go — by trade and by person. Big numbers in travel, pickups, waiting or callbacks are where you're slow.")
ws["A3"] = '="Analysis period: "&AnalysisPeriod&"   ·   Counts ALL logged hours in the period (in-progress jobs included)"'
ws["A3"].font = Font(name=FONT, size=10, bold=True, color=GREEN_T)
NT = 20
TL = lambda k: W("Time Log", k)
ws.column_dimensions["A"].width = 22
section(ws, 4, "1. Hours by Trade × Task", 1, NT + 3)
r0 = 5
hdr(ws, r0, 1, "Trade")
for k in range(NT):
    hdr(ws, r0, 2 + k, None)
    ws.cell(r0, 2 + k).value = f'=IF({settings_cell("Tasks", k)}="","",{settings_cell("Tasks", k)})'
    ws.column_dimensions[GL(2 + k)].width = 11
hdr(ws, r0, NT + 2, "Total Hours")
hdr(ws, r0, NT + 3, "Wrench-Time %", n="Share of this trade's hours spent on productive tasks.")
ws.column_dimensions[GL(NT + 2)].width = 11
ws.column_dimensions[GL(NT + 3)].width = 11
ws.row_dimensions[r0].height = 45
rows_trade = N_TRADES_ROWS + 1
for i in range(rows_trade):
    r = r0 + 1 + i
    lab = settings_cell("Trades", i) if i < N_TRADES_ROWS else None
    out(ws, r, 1, f'=IF({lab}="","",{lab})' if lab else "Overhead", bold=True)
    for k in range(NT):
        Lk = GL(2 + k)
        out(ws, r, 2 + k, f'=IF(OR($A{r}="",{Lk}${r0}=""),"",SUMIFS({TL("Hours")},{TL("ProjectTrade")},$A{r},{TL("Task")},{Lk}${r0},{TL("InPeriod")},1))', '0.0;-0.0;"-"')
    Lt = GL(NT + 2)
    out(ws, r, NT + 2, f'=IF($A{r}="","",SUM(B{r}:{GL(NT + 1)}{r}))', '0.0;-0.0;"-"', bold=True)
    out(ws, r, NT + 3, f'=IF(OR($A{r}="",N({Lt}{r})=0),"",SUMIFS({TL("Hours")},{TL("ProjectTrade")},$A{r},{TL("Productive")},"Yes",{TL("InPeriod")},1)/{Lt}{r})', PCT)
tot1 = r0 + 1 + rows_trade
out(ws, tot1, 1, "TOTAL", bold=True, fill=fill_total)
for k in range(NT + 1):
    L = GL(2 + k)
    out(ws, tot1, 2 + k, f"=SUM({L}{r0 + 1}:{L}{tot1 - 1})", '0.0;-0.0;"-"', bold=True, fill=fill_total)
Lt = GL(NT + 2)
out(ws, tot1, NT + 3, f'=IFERROR(SUMIFS({TL("Hours")},{TL("Productive")},"Yes",{TL("InPeriod")},1)/{Lt}{tot1},"")', PCT, bold=True, fill=fill_total)

s2 = tot1 + 2
section(ws, s2, "2. Share of Each Trade's Hours by Task (row = 100%) — darker = more of the time", 1, NT + 3)
r1 = s2 + 1
hdr(ws, r1, 1, "Trade")
for k in range(NT):
    hdr(ws, r1, 2 + k, None)
    ws.cell(r1, 2 + k).value = f"={GL(2 + k)}${r0}"
ws.row_dimensions[r1].height = 45
for i in range(rows_trade + 1):
    r = r1 + 1 + i
    src = r0 + 1 + i
    out(ws, r, 1, f'=A{src}', bold=True)
    for k in range(NT):
        L = GL(2 + k)
        out(ws, r, 2 + k, f'=IF(OR($A{r}="",N({Lt}{src})=0,{L}${r1}=""),"",{L}{src}/{Lt}{src})', PCT)
ws.conditional_formatting.add(f"B{r1 + 1}:{GL(NT + 1)}{r1 + 1 + rows_trade}",
                              ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="max", end_color="F4B183"))

s3 = r1 + rows_trade + 3
section(ws, s3, "3. Hours by Person × Task", 1, NT + 3)
r2 = s3 + 1
hdr(ws, r2, 1, "Worker")
for k in range(NT):
    hdr(ws, r2, 2 + k, None)
    ws.cell(r2, 2 + k).value = f"={GL(2 + k)}${r0}"
hdr(ws, r2, NT + 2, "Total Hours")
hdr(ws, r2, NT + 3, "Wrench-Time %")
ws.row_dimensions[r2].height = 45
TEAM_L = SHEETS["Team"]["letters"]
for i in range(30):
    r = r2 + 1 + i
    tr_ = FIRST + i
    out(ws, r, 1, f"=IF(Team!${TEAM_L['WorkerID']}${tr_}=\"\",\"\",Team!${TEAM_L['WorkerID']}${tr_}&\" – \"&Team!${TEAM_L['Name']}${tr_})", bold=True)
    for k in range(NT):
        Lk = GL(2 + k)
        out(ws, r, 2 + k, f'=IF(OR($A{r}="",{Lk}${r2}=""),"",SUMIFS({TL("Hours")},{TL("WorkerID")},Team!${TEAM_L["WorkerID"]}${tr_},{TL("Task")},{Lk}${r2},{TL("InPeriod")},1))', '0.0;-0.0;"-"')
    out(ws, r, NT + 2, f'=IF($A{r}="","",SUM(B{r}:{GL(NT + 1)}{r}))', '0.0;-0.0;"-"', bold=True)
    out(ws, r, NT + 3, f'=IF(OR($A{r}="",N({Lt}{r})=0),"",SUMIFS({TL("Hours")},{TL("WorkerID")},Team!${TEAM_L["WorkerID"]}${tr_},{TL("Productive")},"Yes",{TL("InPeriod")},1)/{Lt}{r})', PCT)
ws.freeze_panes = "B6"

# ---------------- Pricing & Hiring
ws = WS["Pricing & Hiring"]
title(ws, "Pricing & Hiring", "What an hour of your time really costs, what you should charge, how you compare to the market — and whether a new hire would pay for themselves.")
for L, w in zip("ABCDEFGHIJKLMN", [44, 16, 16, 16, 16, 16, 14, 14, 14, 14, 11, 16, 40, 4]):
    ws.column_dimensions[L].width = w
EXP = lambda k: W("Expenses", k)
PJ = lambda k: W("Projects", k)

section(ws, 4, "1. Overhead Budget — your fixed business costs (yellow = fill in monthly amount)", 1, 5)
for j, h in enumerate(["Overhead Category", "Monthly Budget $", "Annual Budget $", "Actual, Last 12 Months $", "Actual vs Budget $"]):
    hdr(ws, 5, j + 1, h)
oh_start = 6
cat_col = LIST_POS["ExpCategories"][0]
for i, cname in enumerate(OH_CATS):
    r = oh_start + i
    idx = len(JOB_CATS) + i
    out(ws, r, 1, f"=Settings!${GL(cat_col)}${FIRST + idx}", bold=True)
    c = out(ws, r, 2, None, USD0)
    c.fill, c.font = fill_yellow, f_in
    out(ws, r, 3, f"=N(B{r})*12", USD0)
    out(ws, r, 4, f'=SUMIFS({EXP("Total")},{EXP("Category")},A{r},{EXP("Date")},">="&EDATE(TODAY(),-12),{EXP("Date")},"<="&TODAY())', USD0)
    out(ws, r, 5, f"=D{r}-C{r}", USD0)
oh_tot = oh_start + len(OH_CATS)
out(ws, oh_tot, 1, "TOTAL OVERHEAD", bold=True, fill=fill_total)
for c in (2, 3, 4, 5):
    L = GL(c)
    out(ws, oh_tot, c, f"=SUM({L}{oh_start}:{L}{oh_tot - 1})", USD0, bold=True, fill=fill_total)
ws.cell(oh_start, 13, "Owner draws / distributions are NOT overhead — owners are paid through their hourly rate on the Team tab (that labor is counted in job costs) and profit. Don't put personal spending here.").font = f_sub
ws.cell(oh_start, 13).alignment = wrap_l
ws.merge_cells(start_row=oh_start, start_column=13, end_row=oh_start + 4, end_column=13)

s = oh_tot + 2
section(ws, s, "2. Capacity — how many hours you can sell (yellow = your inputs)", 1, 5)
cap = [("Field workers (people who do jobs, incl. owners)", 2, INT, "CapWorkers", "Count everyone whose hours go on jobs."),
       ("Paid hours per worker per year", 2000, NUM, "CapHours", "40 hrs × 50 weeks = 2,000."),
       ("Share of paid hours logged to jobs (%)", 0.8, PCT, "CapShare", "The rest is admin, estimating, shop, training. Check your real number on the Dashboard once you have data.")]
rr = s + 1
for lbl, v, fmt, nm, nt in cap:
    out(ws, rr, 1, lbl, bold=True)
    c = out(ws, rr, 2, v, fmt)
    c.fill, c.font = fill_yellow, f_in
    ws.cell(rr, 3, nt).font = f_sub
    name(nm, f"'Pricing & Hiring'!$B${rr}")
    rr += 1
out(ws, rr, 1, "Total paid hours per year", bold=True); out(ws, rr, 2, "=CapWorkers*CapHours", NUM); name("CapPaid", f"'Pricing & Hiring'!$B${rr}"); rr += 1
out(ws, rr, 1, "Job (project) hours per year", bold=True); out(ws, rr, 2, "=CapPaid*CapShare", NUM); name("CapProj", f"'Pricing & Hiring'!$B${rr}"); rr += 1
out(ws, rr, 1, "Non-job hours per year (paid, but not billable to a job)", bold=True); out(ws, rr, 2, "=CapPaid-CapProj", NUM); name("CapNonProj", f"'Pricing & Hiring'!$B${rr}"); rr += 1

s = rr + 1
section(ws, s, "3. What One Job Hour Costs You — and What to Charge", 1, 5)
rr = s + 1
TLc = lambda k: W("Time Log", k)
calc = [
    ("Blended labor cost per hour (loaded)", f'=IFERROR(SUMIFS({TLc("LaborCost")},{TLc("InPeriod")},1)/SUMIFS({TLc("Hours")},{TLc("InPeriod")},1),IFERROR(AVERAGEIFS({W("Team","LoadedRate")},{W("Team","Status")},"Active"),0))', USD, "BlendedLaborRate",
     "Average loaded cost of an hour of crew time, from the Time Log (falls back to the Team average until you log hours)."),
    ("Annual overhead (budget above; if blank, actual last 12 months)", f"=IF(C{oh_tot}>0,C{oh_tot},D{oh_tot})", USD0, "AnnualOverhead", ""),
    ("Plus: cost of non-job hours (admin, estimating, shop…)", "=CapNonProj*BlendedLaborRate", USD0, "NonJobLabor",
     "You pay for these hours but can't bill them to a job, so they're overhead too."),
    ("Total overhead to recover", "=AnnualOverhead+NonJobLabor", USD0, "TotalOverhead", ""),
    ("Overhead per job hour", '=IFERROR(TotalOverhead/CapProj,0)', USD, "OverheadPerHour",
     "Every job hour must cover this much overhead before you make a dollar of profit. Used for 'Overhead Share' on each project."),
    ("BREAK-EVEN labor rate per job hour", "=BlendedLaborRate+OverheadPerHour", USD, "BreakEvenRate",
     "Charging less than this per hour of labor loses money, even if the job 'feels' profitable."),
    ("Target net profit % (Settings)", "=TargetNet", PCT, None, ""),
    ("TARGET labor rate per job hour", "=IFERROR(BreakEvenRate/(1-TargetNet),0)", USD, "TargetRate",
     "Break-even ÷ (1 − target net profit %). Minimum you should earn per crew hour on labor. Materials are marked up on top."),
    ("Target gross profit per job hour", "=TargetRate-BlendedLaborRate", USD, "TargetGPPerHour",
     "What each job hour should leave after paying the crew and materials. Projects/trades at or above this are green."),
    ("Reality check — actual overhead per job hour, last 12 months",
     f'=IFERROR((SUMIFS({EXP("Total")},{EXP("CostType")},"Overhead",{EXP("Date")},">="&EDATE(TODAY(),-12))+SUMIFS({TLc("LaborCost")},{TLc("ProjectID")},"OVERHEAD",{TLc("Date")},">="&EDATE(TODAY(),-12)))/SUMIFS({TLc("Hours")},{TLc("ProjectID")},"<>OVERHEAD",{TLc("Date")},">="&EDATE(TODAY(),-12)),"")', USD, None,
     "From your real Expenses + Time Log. If this is far from the planned number above, fix your budget or capacity inputs."),
    ("Reality check — actual gross profit per job hour (analysis period)",
     f'=IFERROR(SUMIFS({PJ("GrossProfit")},{PJ("InPeriod")},1)/SUMIFS({PJ("LaborHours")},{PJ("InPeriod")},1),"")', USD, "ActualGPPerHour", ""),
]
for lbl, f, fmt, nm, nt in calc:
    strong = lbl.startswith(("BREAK", "TARGET"))
    out(ws, rr, 1, lbl, bold=True, fill=fill_total if strong else None)
    out(ws, rr, 2, f, fmt, bold=strong, fill=fill_total if strong else None)
    if nt:
        ws.cell(rr, 3, nt).font = f_sub
    if nm:
        name(nm, f"'Pricing & Hiring'!$B${rr}")
    rr += 1

s = rr + 1
section(ws, s, "4. Rate by Trade vs. the Market", 1, 13)
rh = s + 1
heads = ["Trade", "Our Labor Cost $/hr", "Break-Even Labor Rate $/hr", "Target Labor Rate $/hr", "Our Actual Effective Labor Rate $/hr",
         "Market Hourly Rate — Low", "Market Hourly Rate — Avg", "Market Hourly Rate — High", "# Hourly Rates Logged",
         "Suggested Labor Rate $/hr", "", "Actual GP $/hr", "Position"]
notes4 = {4: "(Revenue − materials, subs & other job costs) ÷ labor hours on completed jobs. What your labor actually earned per hour — directly comparable to a competitor's hourly rate.",
          9: "Your target rate, or the market's average hourly rate if the market pays more. Never below target.",
          5: "Only Market Rates rows with Pricing Basis = Hourly (competitors' posted labor rates), recent only."}
for j, h in enumerate(heads):
    if h:
        hdr(ws, rh, j + 1, h, n=notes4.get(j))
ws.row_dimensions[rh].height = 45
MK = lambda k: W("Market Rates", k)
TA = SHEETS["Trade Analysis"]["letters"]
for i in range(N_TRADES_ROWS):
    r = rh + 1 + i
    lab = settings_cell("Trades", i)
    out(ws, r, 1, f'=IF({lab}="","",{lab})', bold=True)
    g = lambda body: f'=IF($A{r}="","",{body})'
    out(ws, r, 2, g(f'IFERROR(SUMIFS({TLc("LaborCost")},{TLc("ProjectTrade")},$A{r},{TLc("InPeriod")},1)/SUMIFS({TLc("Hours")},{TLc("ProjectTrade")},$A{r},{TLc("InPeriod")},1),BlendedLaborRate)'), USD)
    out(ws, r, 3, g(f"B{r}+OverheadPerHour"), USD)
    out(ws, r, 4, g(f"IFERROR(C{r}/(1-TargetNet),0)"), USD)
    tp = f'{PJ("Trade")},$A{r},{PJ("InPeriod")},1'
    out(ws, r, 5, g(f'IFERROR((SUMIFS({PJ("Revenue")},{tp})-SUMIFS({PJ("DirectCost")},{tp})+SUMIFS({PJ("LaborCost")},{tp}))/SUMIFS({PJ("LaborHours")},{tp}),"")'), USD)
    mb = f'{MK("Trade")},$A{r},{MK("Recent")},1,{MK("Basis")},"Hourly"'
    out(ws, r, 6, g(f'IF(I{r}=0,"",_xlfn.MINIFS({MK("ImpliedHourly")},{mb}))'), USD)
    out(ws, r, 7, g(f'IFERROR(AVERAGEIFS({MK("ImpliedHourly")},{mb}),"")'), USD)
    out(ws, r, 8, g(f'IF(I{r}=0,"",_xlfn.MAXIFS({MK("ImpliedHourly")},{mb}))'), USD)
    out(ws, r, 9, g(f'COUNTIFS({mb},{MK("ImpliedHourly")},">0")'), INT)
    out(ws, r, 10, g(f'IF(G{r}="",D{r},MAX(D{r},G{r}))'), USD, bold=True)
    out(ws, r, 12, g(f"'Trade Analysis'!{TA['GPHr']}{FIRST + i}"), USD)
    out(ws, r, 13, g(f'IF(AND(E{r}="",G{r}=""),"Add jobs & market prices",IF(AND(E{r}<>"",E{r}<C{r}),"⚠ Earning below break-even",IF(G{r}="","Add competitor hourly rates on Market Rates",IF(G{r}>=D{r},"Market avg ~$"&TEXT(G{r},"0")&"/hr is above your target — charge at least that","Market hourly rates are below your target — sell on quality/speed, or cut cost"))))'))
verdict_cf(ws, f"M{rh + 1}:M{rh + N_TRADES_ROWS}", f"M{rh + 1}")

s = rh + N_TRADES_ROWS + 2
section(ws, s, "5. Hiring Calculator — would one more person pay for themselves? (yellow = inputs)", 1, 5)
rr = s + 1
hire_in = [("New hire pay rate ($/hr)", 25, USD, "HPay", ""),
           ("Burden % (payroll tax, workers' comp, benefits)", "=BurdenDefault", PCT, "HBurden", "Defaults to Settings. Use 0% for a 1099 sub."),
           ("Paid hours per year", 2000, NUM, "HHours", ""),
           ("Share of their hours on jobs (%)", "=CapShare", PCT, "HShare", ""),
           ("Productivity vs. your current crew (%)", 0.7, PCT, "HProd", "A new helper is rarely as fast as you. 60–80% in year one is realistic."),
           ("Added yearly overhead for this hire ($)", 0, USD0, "HOH", "Extra truck, tools, phone, insurance, uniforms…"),
           ("One-time startup cost ($)", 0, USD0, "HOnce", "Recruiting, training time, tool kit.")]
for lbl, v, fmt, nm, nt in hire_in:
    out(ws, rr, 1, lbl, bold=True)
    c = out(ws, rr, 2, v, fmt)
    c.fill, c.font = fill_yellow, f_in
    if nt:
        ws.cell(rr, 3, nt).font = f_sub
    name(nm, f"'Pricing & Hiring'!$B${rr}")
    rr += 1
hire_out = [
    ("Contribution per job hour, before labor (from your data)",
     f'=IFERROR((SUMIFS({PJ("Revenue")},{PJ("InPeriod")},1)-(SUMIFS({PJ("DirectCost")},{PJ("InPeriod")},1)-SUMIFS({PJ("LaborCost")},{PJ("InPeriod")},1)))/SUMIFS({PJ("LaborHours")},{PJ("InPeriod")},1),TargetRate)', USD, "HContrib",
     "(Revenue − materials/subs/other job costs) ÷ labor hours on completed jobs. Uses your target rate until you have data."),
    ("Their yearly cost (loaded)", "=HPay*(1+HBurden)*HHours", USD0, "HCost", ""),
    ("Job hours they'd work per year", "=HHours*HShare", NUM, "HJobHrs", ""),
    ("Contribution they'd generate per year", "=HJobHrs*HContrib*HProd", USD0, "HGen", ""),
    ("NET YEARLY GAIN (loss) from this hire", "=HGen-HCost-HOH", USD0, "HNet", ""),
    ("Job hours needed just to break even", '=IFERROR((HCost+HOH)/(HContrib*HProd),"")', NUM, None,
     "You need at least this much extra work lined up for them."),
    ("Extra revenue you must sell per year to keep them busy", "=HJobHrs*HProd*IFERROR(SUMIFS(" + PJ("Revenue") + "," + PJ("InPeriod") + ",1)/SUMIFS(" + PJ("LaborHours") + "," + PJ("InPeriod") + ",1),TargetRate)", USD0, None, ""),
    ("Months to pay back startup cost", '=IF(HNet<=0,"Never at these numbers",IF(HOnce=0,0,HOnce/(HNet/12)))', "0.0", None, ""),
    ("Verdict", '=IF(HNet>0,"✔ Pays for itself — IF you have the work to fill "&TEXT(HJobHrs,"#,##0")&" job hours a year","✘ Doesn\'t pay at these numbers — raise prices, raise productivity, or wait for more work")', None, None, ""),
]
for lbl, f, fmt, nm, nt in hire_out:
    strong = lbl.startswith(("NET", "Verdict"))
    out(ws, rr, 1, lbl, bold=True, fill=fill_total if strong else None)
    out(ws, rr, 2, f, fmt, bold=strong, fill=fill_total if strong else None)
    if nt:
        ws.cell(rr, 3, nt).font = f_sub
    if nm:
        name(nm, f"'Pricing & Hiring'!$B${rr}")
    if lbl == "Verdict":
        ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=8)
        verdict_cf(ws, f"B{rr}", f"B{rr}")
        ws.conditional_formatting.add(f"B{rr}", FormulaRule(formula=[f'LEFT(B{rr},1)="✔"'], fill=GRN_F))
        ws.conditional_formatting.add(f"B{rr}", FormulaRule(formula=[f'LEFT(B{rr},1)="✘"'], fill=RED_F))
    rr += 1

# ---------------- Marketing
ws = WS["Marketing"]
title(ws, "Marketing & Lead Sources", "Where your best jobs come from and what each lead source really costs. Log marketing spend on Expenses with a Marketing Channel.")
ws["A3"] = '="Analysis period: "&AnalysisPeriod'
ws["A3"].font = Font(name=FONT, size=10, bold=True, color=GREEN_T)
QT = lambda k: W("Quotes", k)
mk_cols = ["Lead Source", "Leads (Quotes)", "Won", "Win Rate %", "Completed Projects", "Revenue $", "Gross Profit $",
           "Gross Margin %", "Avg Job Revenue $", "Marketing Spend $", "Cost per Lead $", "Cost per Won Job $",
           "Gross Profit After Marketing $", "Marketing ROI", "Share of Revenue", "Verdict"]
mk_w = [30, 9, 8, 9, 10, 12, 12, 9, 11, 11, 10, 11, 13, 9, 9, 32]
mk_notes = {13: "(Gross profit − spend) ÷ spend. 3.0x = every $1 of ads brought back $3 of profit after paying for the ads."}
for j, h in enumerate(mk_cols):
    hdr(ws, HDR_ROW, j + 1, h, n=mk_notes.get(j))
    ws.column_dimensions[GL(j + 1)].width = mk_w[j]
ws.row_dimensions[HDR_ROW].height = 45
NLS = 25
for i in range(NLS):
    r = FIRST + i
    lab = settings_cell("LeadSources", i)
    out(ws, r, 1, f'=IF({lab}="","",{lab})', bold=True)
    g = lambda body, rr_=r: f'=IF($A{rr_}="","",{body})'
    qb = f'{QT("LeadSource")},$A{r},{QT("InPeriod")},1'
    pb = f'{PJ("LeadSource")},$A{r},{PJ("InPeriod")},1'
    out(ws, r, 2, g(f"COUNTIFS({qb})"), INT)
    out(ws, r, 3, g(f'COUNTIFS({qb},{QT("Status")},"Won")'), INT)
    out(ws, r, 4, g(f'IFERROR(C{r}/(C{r}+COUNTIFS({qb},{QT("Status")},"Lost")+COUNTIFS({qb},{QT("Status")},"Expired")),"")'), PCT)
    out(ws, r, 5, g(f"COUNTIFS({pb})"), INT)
    out(ws, r, 6, g(f'SUMIFS({PJ("Revenue")},{pb})'), USD0)
    out(ws, r, 7, g(f'SUMIFS({PJ("GrossProfit")},{pb})'), USD0)
    out(ws, r, 8, g(f'IF(F{r}=0,"",G{r}/F{r})'), PCT)
    out(ws, r, 9, g(f'IF(E{r}=0,"",F{r}/E{r})'), USD0)
    out(ws, r, 10, g(f'SUMIFS({EXP("Total")},{EXP("Channel")},$A{r},{EXP("InPeriod")},1)'), USD0)
    out(ws, r, 11, g(f'IF(OR(B{r}=0,J{r}=0),"",J{r}/B{r})'), USD0)
    out(ws, r, 12, g(f'IF(OR(C{r}=0,J{r}=0),"",J{r}/C{r})'), USD0)
    out(ws, r, 13, g(f"G{r}-J{r}"), USD0)
    out(ws, r, 14, g(f'IF(J{r}=0,"",M{r}/J{r})'), MULT)
    out(ws, r, 15, g(f'IFERROR(F{r}/SUM($F${FIRST}:$F${FIRST + NLS - 1}),"")'), PCT)
    out(ws, r, 16, g(f'IF(AND(B{r}=0,J{r}=0),"",IF(J{r}=0,IF(G{r}>0,"Free channel — nurture it",""),IF(M{r}<0,"⚠ Costs more than it earns",IF(N(N{r})>=3,"★ Strong — consider more spend","OK — keep testing"))))'))
mt = FIRST + NLS
out(ws, mt, 1, "TOTAL", bold=True, fill=fill_total)
for c in (2, 3, 5, 6, 7, 10, 13):
    L = GL(c)
    out(ws, mt, c, f"=SUM({L}{FIRST}:{L}{mt - 1})", INT if c in (2, 3, 5) else USD0, bold=True, fill=fill_total)
out(ws, mt, 8, f'=IFERROR(G{mt}/F{mt},"")', PCT, bold=True, fill=fill_total)
out(ws, mt, 12, f'=IFERROR(J{mt}/C{mt},"")', USD0, bold=True, fill=fill_total)
out(ws, mt, 14, f'=IFERROR(M{mt}/J{mt},"")', MULT, bold=True, fill=fill_total)
verdict_cf(ws, f"P{FIRST}:P{mt - 1}", f"P{FIRST}")
ws.freeze_panes = "B5"

# ---------------- Clients & Areas
ws = WS["Clients & Areas"]
title(ws, "Clients & Areas", "Which kinds of clients and which areas are worth the most — and who your best clients are.")
ws["A3"] = '="Analysis period: "&AnalysisPeriod&"   (Tiers and Top 10 are all-time)"'
ws["A3"].font = Font(name=FONT, size=10, bold=True, color=GREEN_T)
for L, w in zip("ABCDEFGHIJKL", [30, 12, 12, 13, 13, 10, 11, 11, 12, 12, 12, 30]):
    ws.column_dimensions[L].width = w
CL = lambda k: W("Clients", k)
PR = lambda k: W("Properties", k)
section(ws, 4, "1. By Client Type", 1, 12)
h1 = ["Client Type", "# Clients", "Completed Projects", "Revenue $", "Gross Profit $", "Gross Margin %", "Labor Hours",
      "GP / Labor Hr", "Avg Job Revenue $", "Avg Days to Get Paid", "Avg Satisfaction", "Repeat Clients"]
for j, h in enumerate(h1):
    hdr(ws, 5, j + 1, h)
ws.row_dimensions[5].height = 40
for i in range(10):
    r = 6 + i
    lab = settings_cell("ClientTypes", i)
    out(ws, r, 1, f'=IF({lab}="","",{lab})', bold=True)
    g = lambda body, rr_=r: f'=IF($A{rr_}="","",{body})'
    pb = f'{PJ("ClientType")},$A{r},{PJ("InPeriod")},1'
    out(ws, r, 2, g(f'COUNTIFS({CL("ClientType")},$A{r})'), INT)
    out(ws, r, 3, g(f"COUNTIFS({pb})"), INT)
    out(ws, r, 4, g(f'SUMIFS({PJ("Revenue")},{pb})'), USD0)
    out(ws, r, 5, g(f'SUMIFS({PJ("GrossProfit")},{pb})'), USD0)
    out(ws, r, 6, g(f'IF(D{r}=0,"",E{r}/D{r})'), PCT)
    out(ws, r, 7, g(f'SUMIFS({PJ("LaborHours")},{pb})'), NUM)
    out(ws, r, 8, g(f'IF(G{r}=0,"",E{r}/G{r})'), USD)
    out(ws, r, 9, g(f'IF(C{r}=0,"",D{r}/C{r})'), USD0)
    out(ws, r, 10, g(f'IFERROR(AVERAGEIFS({PJ("DaysToCollect")},{pb}),"")'), "0.0")
    out(ws, r, 11, g(f'IFERROR(AVERAGEIFS({PJ("Satisfaction")},{pb}),"")'), "0.0")
    out(ws, r, 12, g(f'COUNTIFS({CL("ClientType")},$A{r},{CL("Repeat")},"Yes")'), INT)

section(ws, 17, "2. By Travel Zone (distance from your base)", 1, 12)
h2 = ["Travel Zone", "# Properties", "Completed Projects", "Revenue $", "Gross Profit $", "Gross Margin %", "Labor Hours",
      "GP / Labor Hr", "Travel/Pickup/Wait % of Hours", "Avg Drive (min, one way)", "", "Verdict"]
for j, h in enumerate(h2):
    if h:
        hdr(ws, 18, j + 1, h)
ws.row_dimensions[18].height = 40
for i in range(4):
    r = 19 + i
    out(ws, r, 1, f"=INDEX(List_Zones,{i + 1})", bold=True)
    pb = f'{PJ("TravelZone")},$A{r},{PJ("InPeriod")},1'
    out(ws, r, 2, f'=COUNTIFS({PR("TravelZone")},$A{r})', INT)
    out(ws, r, 3, f"=COUNTIFS({pb})", INT)
    out(ws, r, 4, f'=SUMIFS({PJ("Revenue")},{pb})', USD0)
    out(ws, r, 5, f'=SUMIFS({PJ("GrossProfit")},{pb})', USD0)
    out(ws, r, 6, f'=IF(D{r}=0,"",E{r}/D{r})', PCT)
    out(ws, r, 7, f'=SUMIFS({PJ("LaborHours")},{pb})', NUM)
    out(ws, r, 8, f'=IF(G{r}=0,"",E{r}/G{r})', USD)
    out(ws, r, 9, f'=IF(G{r}=0,"",SUMIFS({PJ("NonProdHours")},{pb})/G{r})', PCT)
    out(ws, r, 10, f'=IFERROR(AVERAGEIFS({PR("DriveMin")},{PR("TravelZone")},$A{r}),"")', "0")
    out(ws, r, 12, f'=IF(C{r}=0,"No data yet",IF(H{r}="","Log hours to evaluate",IF(H{r}<OverheadPerHour,"⚠ Far jobs losing money — add a trip charge",IF(H{r}>=TargetGPPerHour,"★ Worth the drive","OK"))))')
verdict_cf(ws, "L19:L22", "L19")

section(ws, 24, "3. Client Tiers (all-time; thresholds on Settings)", 1, 12)
for j, h in enumerate(["Tier", "# Clients", "Lifetime Revenue $", "Lifetime Gross Profit $", "Share of All Gross Profit"]):
    hdr(ws, 25, j + 1, h)
for i, t in enumerate(["A", "B", "C"]):
    r = 26 + i
    out(ws, r, 1, t, bold=True)
    out(ws, r, 2, f'=COUNTIFS({CL("Tier")},A{r})', INT)
    out(ws, r, 3, f'=SUMIFS({CL("LTRev")},{CL("Tier")},A{r})', USD0)
    out(ws, r, 4, f'=SUMIFS({CL("LTGP")},{CL("Tier")},A{r})', USD0)
    out(ws, r, 5, f'=IFERROR(D{r}/SUM($D$26:$D$28),"")', PCT)

section(ws, 30, "4. Top 10 Clients by Lifetime Gross Profit", 1, 12)
h4 = ["Rank", "Client ID", "Client Name", "Client Type", "Completed Projects", "Lifetime Revenue $", "Lifetime Gross Profit $",
      "Avg Margin %", "Last Project", "Days Since", "Referrals Made", "Tier"]
for j, h in enumerate(h4):
    hdr(ws, 31, j + 1, h)
ws.row_dimensions[31].height = 40
fields = [None, "ClientID", "DisplayName", "ClientType", "Completed", "LTRev", "LTGP", "AvgGM", "LastProject", "DaysSince", "Referrals", "Tier"]
fmts = [INT, None, None, None, INT, USD0, USD0, PCT, DATE, INT, INT, None]
for i in range(10):
    r = 32 + i
    key = f"LARGE({CL('SortKey')},{i + 1})"
    mrow = f"MATCH({key},{CL('SortKey')},0)"
    out(ws, r, 1, i + 1, INT, bold=True)
    for j in range(1, 12):
        out(ws, r, j + 1, f'=IFERROR(IF(INDEX({CL("LTGP")},{mrow})<=0,"",INDEX({CL(fields[j])},{mrow})),"")', fmts[j])

# ---------------- Monthly Trend
ws = WS["Monthly Trend"]
title(ws, "Monthly Trend", "Month-by-month money, hours and sales for the year you pick. Revenue is counted in the month a project is completed.")
ws["A4"] = "Year:"
ws["A4"].font = f_bold
ws["B4"] = 2026
ws["B4"].font, ws["B4"].fill, ws["B4"].border = f_in, fill_yellow, box
name("TrendYear", "'Monthly Trend'!$B$4")
mh = ["Month", "Revenue $", "Direct Cost $", "Gross Profit $", "Gross Margin %", "Overhead $ (incl. non-job labor)",
      "Net Profit $", "Cash Collected $", "Job Hours", "Total Hours Paid", "Job-Hour Share %", "Quotes Received",
      "Quotes Won", "Win Rate %", "New Clients", "Projects Completed", "Marketing Spend $"]
mw = [11, 12, 12, 12, 9, 13, 12, 12, 9, 9, 9, 9, 8, 8, 8, 9, 11]
for j, h in enumerate(mh):
    hdr(ws, 6, j + 1, h)
    ws.column_dimensions[GL(j + 1)].width = mw[j]
ws.row_dimensions[6].height = 45
IV = lambda k: W("Invoices", k)
for m in range(12):
    r = 7 + m
    out(ws, r, 1, f"=DATE(TrendYear,{m + 1},1)", "mmm yyyy", bold=True)
    rng = lambda col: f'{col},">="&$A{r},{col},"<="&EOMONTH($A{r},0)'
    pc = f'{PJ("Status")},"Completed",{rng(PJ("ActualEnd"))}'
    out(ws, r, 2, f'=SUMIFS({PJ("Revenue")},{pc})', USD0)
    out(ws, r, 3, f'=SUMIFS({PJ("DirectCost")},{pc})', USD0)
    out(ws, r, 4, f"=B{r}-C{r}", USD0)
    out(ws, r, 5, f'=IF(B{r}=0,"",D{r}/B{r})', PCT)
    out(ws, r, 6, f'=SUMIFS({EXP("Total")},{EXP("CostType")},"Overhead",{rng(EXP("Date"))})+SUMIFS({TLc("LaborCost")},{TLc("ProjectID")},"OVERHEAD",{rng(TLc("Date"))})', USD0)
    out(ws, r, 7, f"=D{r}-F{r}", USD0)
    out(ws, r, 8, f'=SUMIFS({IV("AmountPaid")},{rng(IV("DatePaid"))})', USD0)
    out(ws, r, 9, f'=SUMIFS({TLc("Hours")},{TLc("ProjectID")},"<>OVERHEAD",{rng(TLc("Date"))})', '0.0;-0.0;"-"')
    out(ws, r, 10, f'=SUMIFS({TLc("Hours")},{rng(TLc("Date"))})', '0.0;-0.0;"-"')
    out(ws, r, 11, f'=IF(J{r}=0,"",I{r}/J{r})', PCT)
    out(ws, r, 12, f'=COUNTIFS({rng(QT("DateReceived"))})', INT)
    out(ws, r, 13, f'=COUNTIFS({QT("Status")},"Won",{rng(QT("DecisionDate"))})', INT)
    out(ws, r, 14, f'=IFERROR(M{r}/(M{r}+COUNTIFS({QT("Status")},"Lost",{rng(QT("DecisionDate"))})+COUNTIFS({QT("Status")},"Expired",{rng(QT("DecisionDate"))})),"")', PCT)
    out(ws, r, 15, f'=COUNTIFS({rng(CL("FirstContact"))})', INT)
    out(ws, r, 16, f'=COUNTIFS({pc})', INT)
    out(ws, r, 17, f'=SUMIFS({EXP("Total")},{EXP("Category")},"Marketing & Advertising",{rng(EXP("Date"))})', USD0)
out(ws, 19, 1, "TOTAL", bold=True, fill=fill_total)
for c in range(2, 18):
    L = GL(c)
    if c in (5, 11, 14):
        f = {5: '=IF(B19=0,"",D19/B19)', 11: '=IF(J19=0,"",I19/J19)',
             14: f'=IFERROR(M19/(M19+COUNTIFS({QT("Status")},"Lost",{QT("DecisionDate")},">="&$A$7,{QT("DecisionDate")},"<="&EOMONTH($A$18,0))+COUNTIFS({QT("Status")},"Expired",{QT("DecisionDate")},">="&$A$7,{QT("DecisionDate")},"<="&EOMONTH($A$18,0))),"")'}[c]
        out(ws, 19, c, f, PCT, bold=True, fill=fill_total)
    else:
        out(ws, 19, c, f"=SUM({L}7:{L}18)", INT if c in (12, 13, 15, 16) else ('0.0;-0.0;"-"' if c in (9, 10) else USD0), bold=True, fill=fill_total)
ch = BarChart()
ch.type, ch.title, ch.height, ch.width = "col", "Revenue, Gross Profit & Net Profit by Month", 8, 22
ch.add_data(Reference(ws, min_col=2, min_row=6, max_row=18), titles_from_data=True)
ch.add_data(Reference(ws, min_col=4, min_row=6, max_row=18), titles_from_data=True)
ch.add_data(Reference(ws, min_col=7, min_row=6, max_row=18), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=7, max_row=18))
ch.y_axis.numFmt = "$#,##0"
ch.y_axis.delete = False
ch.x_axis.delete = False
ch.x_axis.number_format = "mmm"
ws.add_chart(ch, "A22")
ws.freeze_panes = "B7"

# ---------------- Dashboard
ws = WS["Dashboard"]
ws["B1"] = "=BizName&\" — Business Dashboard\""
ws["B1"].font = f_title
ws["B2"] = "The numbers that run the business. Pick the period below — every analysis tab follows it."
ws["B2"].font = f_sub
ws.column_dimensions["A"].width = 2
for L in "BCDEFGHI":
    ws.column_dimensions[L].width = 15
ws.column_dimensions["J"].width = 3
ws["B4"] = "Analysis Period:"
ws["B4"].font = f_bold
ws["C4"] = "All Time"
ws["C4"].font, ws["C4"].fill, ws["C4"].border = Font(name=FONT, size=11, bold=True, color="0000FF"), fill_yellow, box
name("AnalysisPeriod", "Dashboard!$C$4")
dvp = DataValidation(type="list", formula1='"All Time,2024,2025,2026,2027,2028,2029,2030"', allow_blank=False)
ws.add_data_validation(dvp)
dvp.add("C4")
ws["D4"] = "← pick All Time or a year"
ws["D4"].font = f_sub

tile_lbl = Font(name=FONT, size=9, bold=True, color="595959")
tile_val = Font(name=FONT, size=16, bold=True, color=NAVY)
tile_fill = PatternFill("solid", fgColor="F3F6FA")
def tile(r, k, label, f, fmt, n=None):
    c1 = 2 + 2 * k
    ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c1 + 1)
    ws.merge_cells(start_row=r + 1, start_column=c1, end_row=r + 1, end_column=c1 + 1)
    a = ws.cell(r, c1, label)
    a.font, a.alignment, a.fill = tile_lbl, Alignment(horizontal="center", wrap_text=True), tile_fill
    b = ws.cell(r + 1, c1, f)
    b.font, b.alignment, b.fill = tile_val, Alignment(horizontal="center", vertical="center"), tile_fill
    b.number_format = fmt
    for cc in (c1, c1 + 1):
        ws.cell(r, cc).fill = tile_fill
        ws.cell(r + 1, cc).fill = tile_fill
        ws.cell(r, cc).border = Border(top=thin, left=thin if cc == c1 else None, right=thin if cc == c1 + 1 else None)
        ws.cell(r + 1, cc).border = Border(bottom=thin, left=thin if cc == c1 else None, right=thin if cc == c1 + 1 else None)
    ws.row_dimensions[r].height = 26
    ws.row_dimensions[r + 1].height = 30
    if n:
        note(a, n)

pin = f'{PJ("InPeriod")},1'
comp = f"COUNTIFS({pin})"
REV = f'SUMIFS({PJ("Revenue")},{pin})'
GP = f'SUMIFS({PJ("GrossProfit")},{pin})'
OH = f'(SUMIFS({EXP("Total")},{EXP("CostType")},"Overhead",{EXP("InPeriod")},1)+SUMIFS({TLc("LaborCost")},{TLc("ProjectID")},"OVERHEAD",{TLc("InPeriod")},1))'
PH = f'SUMIFS({PJ("LaborHours")},{pin})'
qin = f'{QT("InPeriod")},1'
won = f'COUNTIFS({qin},{QT("Status")},"Won")'
decided = f'({won}+COUNTIFS({qin},{QT("Status")},"Lost")+COUNTIFS({qin},{QT("Status")},"Expired"))'
TLH = f'SUMIFS({TLc("Hours")},{TLc("InPeriod")},1)'
D0 = '$#,##0;($#,##0);"$0"'
P0 = '0.0%'
sections = [
    ("MONEY", [
        ("Revenue", f"={REV}", D0, "Completed projects in the period (contract + approved change orders)."),
        ("Gross Profit", f"={GP}", D0, "Revenue − direct job costs (labor, materials, subs, permits, disposal, other)."),
        ("Gross Margin", f'=IFERROR({GP}/{REV},"")', P0, None),
        ("Overhead Spent", f"={OH}", D0, "Overhead expenses + labor cost of non-job (OVERHEAD) hours in the period."),
        ("Net Profit", f"={GP}-{OH}", D0, "Gross profit − overhead spent."),
        ("Net Margin", f'=IFERROR(({GP}-{OH})/{REV},"")', P0, None),
        ("Cash Collected", f'=SUMIFS({IV("AmountPaid")},{IV("PaidInPeriod")},1)', D0, "Payments received in the period."),
        ("Owed to You Now", f'=SUMIFS({IV("Balance")},{IV("Status")},"<>Paid")', D0, "Unpaid balance on all invoices (any period)."),
    ]),
    ("TIME & PRODUCTIVITY", [
        ("Gross Profit / Job Hour", f'=IFERROR({GP}/{PH},"")', '$#,##0.00', "THE number. Compare with your target on Pricing & Hiring."),
        ("Target GP / Job Hour", "=TargetGPPerHour", '$#,##0.00', "From Pricing & Hiring."),
        ("Revenue / Job Hour", f'=IFERROR({REV}/{PH},"")', '$#,##0.00', None),
        ("Hours Paid (all)", f"={TLH}", '#,##0', None),
        ("Wrench-Time %", f'=IFERROR(SUMIFS({TLc("Hours")},{TLc("Productive")},"Yes",{TLc("InPeriod")},1)/{TLH},"")', P0, "Share of paid hours spent actually doing the work."),
        ("Hours vs Estimate", f'=IFERROR(SUMIFS({PJ("LaborHours")},{pin},{PJ("EstHours")},">0")/SUMIFS({PJ("EstHours")},{pin},{PJ("EstHours")},">0")-1,"")', P0, "+ = jobs take longer than you estimate."),
        ("On-Time Finish %", f'=IFERROR(COUNTIFS({pin},{PJ("OnTime")},"Yes")/(COUNTIFS({pin},{PJ("OnTime")},"Yes")+COUNTIFS({pin},{PJ("OnTime")},"No")),"")', P0, None),
        ("Travel / Pickup / Wait Hrs", f'=SUMIFS({TLc("Hours")},{TLc("Task")},"Travel",{TLc("InPeriod")},1)+SUMIFS({TLc("Hours")},{TLc("Task")},"Material Pickup",{TLc("InPeriod")},1)+SUMIFS({TLc("Hours")},{TLc("Task")},"Waiting / Delay",{TLc("InPeriod")},1)', '#,##0.0', "Paid hours that produce nothing — first place to find time."),
    ]),
    ("SALES", [
        ("Quotes Given", f"=COUNTIFS({qin})", "#,##0", None),
        ("Win Rate", f'=IFERROR({won}/{decided},"")', P0, "Won ÷ (Won + Lost + Expired)."),
        ("Avg Days to Send Quote", f'=IFERROR(AVERAGEIFS({QT("DaysToQuote")},{qin}),"")', "0.0", None),
        ("Projects Completed", f"={comp}", "#,##0", None),
        ("Avg Job Size", f'=IFERROR({REV}/{comp},"")', D0, None),
        ("Repeat + Referral Revenue", f'=IFERROR((SUMIFS({PJ("Revenue")},{pin},{PJ("LeadSource")},"Repeat Client")+SUMIFS({PJ("Revenue")},{pin},{PJ("LeadSource")},"Referral*"))/{REV},"")', P0, "Share of revenue that came from people who already know you. Higher = cheaper growth."),
        ("Hours Spent on Lost Quotes", f'=SUMIFS({QT("EstimatingHrs")},{qin},{QT("Status")},"Lost")+SUMIFS({QT("EstimatingHrs")},{qin},{QT("Status")},"Expired")', "#,##0.0", None),
        ("Marketing $ per Won Job", f'=IFERROR(SUMIFS({EXP("Total")},{EXP("Category")},"Marketing & Advertising",{EXP("InPeriod")},1)/{won},"")', D0, None),
    ]),
    ("QUALITY & CLIENTS", [
        ("Callbacks per Project", f'=IFERROR(SUMIFS({PJ("Callbacks")},{pin})/{comp},"")', "0.00", None),
        ("Avg Satisfaction (1-10)", f'=IFERROR(AVERAGEIFS({PJ("Satisfaction")},{pin}),"")', "0.0", None),
        ("Testimonials Collected", f'=IFERROR(COUNTIFS({pin},{PJ("Testimonial")},"Yes")/{comp},"")', P0, "% of completed jobs. Ask on every job you're proud of."),
        ("Online Reviews Left", f'=IFERROR(COUNTIFS({pin},{PJ("Review")},"Yes")/{comp},"")', P0, None),
        ("Inspection 1st-Pass %", f'=IFERROR(COUNTIFS({pin},{PJ("Inspection")},"Yes")/(COUNTIFS({pin},{PJ("Inspection")},"Yes")+COUNTIFS({pin},{PJ("Inspection")},"No")),"")', P0, None),
        ("Avg Days to Get Paid", f'=IFERROR(AVERAGEIFS({PJ("DaysToCollect")},{pin}),"")', "0.0", "Days from finishing the job to being paid in full."),
        ("Overdue Invoices", f'=SUMIFS({IV("Balance")},{IV("Status")},"OVERDUE")', D0, None),
        ("Clients Served (all time)", f'=COUNTIFS({CL("Completed")},">0")', "#,##0", None),
    ]),
]
r = 6
for sname_, tiles in sections:
    section(ws, r, sname_, 2, 9)
    r += 1
    for k in range(0, len(tiles), 4):
        for j, t in enumerate(tiles[k:k + 4]):
            tile(r, j, *t)
        r += 3

section(ws, r, "TRADE SCOREBOARD  (details on Trade Analysis)", 2, 9)
r += 1
sb = ["Trade", "Projects", "Revenue", "Gross Profit", "Margin %", "GP / Hour", "Rank", "Verdict"]
for j, h in enumerate(sb):
    hdr(ws, r, 2 + j, h)
sb_first = r + 1
tacols = ["Label", "Count", "Rev", "GP", "GM", "GPHr", "Rank", "Verdict"]
sbf = [None, INT, USD0, USD0, PCT, USD, INT, None]
for i in range(N_TRADES_ROWS):
    rr = sb_first + i
    for j, k in enumerate(tacols):
        out(ws, rr, 2 + j, f"='Trade Analysis'!{TA[k]}{FIRST + i}", sbf[j], bold=(j == 0))
verdict_cf(ws, f"I{sb_first}:I{sb_first + N_TRADES_ROWS - 1}", f"I{sb_first}")
r = sb_first + N_TRADES_ROWS + 1

section(ws, r, "NEEDS ATTENTION", 2, 9)
r += 1
alerts = [
    ("Overdue invoices", f'=COUNTIFS({IV("Status")},"OVERDUE")'),
    ("Open quotes waiting 7+ days (follow up!)", f'=COUNTIFS({QT("DaysWaiting")},">=7")'),
    ("Projects in progress", f'=COUNTIFS({PJ("Status")},"In Progress")'),
    ("Completed jobs with a balance owed", f'=COUNTIFS({PJ("Status")},"Completed",{PJ("Balance")},">0.004")'),
    ("Team licenses expired / expiring in 60 days", f'=COUNTIFS({W("Team","LicenseFlag")},"EXPIRED")+COUNTIFS({W("Team","LicenseFlag")},"Renew soon")'),
    ("Vendor insurance certs expired / expiring", f'=COUNTIFS({W("Vendors","COIFlag")},"EXPIRED")+COUNTIFS({W("Vendors","COIFlag")},"Renew soon")'),
    ("Properties with future-work flags (call them!)", f'=COUNTIFS({PR("Opportunity")},"•*")'),
    ("Data problems on Time Log + Expenses", f'=COUNTIFS({TLc("Check")},"⚠*")+COUNTIFS({EXP("Check")},"⚠*")'),
]
for lbl, f in alerts:
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    ws.cell(r, 2, lbl).font = f_body
    c = ws.cell(r, 6, f)
    c.font, c.number_format, c.alignment = f_bold, "0", Alignment(horizontal="center")
    ws.conditional_formatting.add(f"F{r}", FormulaRule(formula=[f"F{r}>0"], fill=AMB_F))
    for cc in range(2, 7):
        ws.cell(r, cc).border = Border(bottom=thin)
    r += 1

chart = BarChart()
chart.type, chart.title, chart.height, chart.width = "bar", "Gross Profit per Labor Hour, by Trade", 9, 16
gphr_idx = [i for i, c in enumerate(trade_cols) if c["key"] == "GPHr"][0] + 1
chart.add_data(Reference(WS["Trade Analysis"], min_col=gphr_idx, min_row=HDR_ROW, max_row=FIRST + len(TRADES) - 1), titles_from_data=True)
chart.set_categories(Reference(WS["Trade Analysis"], min_col=1, min_row=FIRST, max_row=FIRST + len(TRADES) - 1))
chart.legend = None
chart.x_axis.delete = False
chart.y_axis.delete = False
chart.y_axis.numFmt = "$#,##0"
ws.add_chart(chart, "K6")

# ---------------- WhatsApp Log (Google Sheets version only)
if GSHEETS:
    ws = WS["WhatsApp Log"]
    title(ws, "WhatsApp Log", "Every message the WhatsApp bot received and exactly what it wrote. Written by the bot — don't edit. Reply UNDO in WhatsApp to reverse the last entry.")
    wl = [("Received", 17, "mm/dd/yyyy h:mm AM/PM"), ("From (phone)", 14, TXT), ("Worker ID", 8, TXT), ("Message", 50, None),
          ("Photo?", 7, None), ("What the bot wrote", 60, None), ("Cells written (for undo)", 30, None),
          ("Status", 10, None), ("Bot reply", 50, None), ("Message ID", 18, TXT)]
    for j, (h, w, fmt) in enumerate(wl):
        hdr(ws, HDR_ROW, j + 1, h, fill=fill_auto_h)
        ws.column_dimensions[GL(j + 1)].width = w
        if fmt:
            for r in range(FIRST, FIRST + 200):
                ws.cell(r, j + 1).number_format = fmt
    ws.row_dimensions[HDR_ROW].height = 32
    ws.freeze_panes = "A5"

# ================================================================ START HERE
ws = WS["Start Here"]
ws.column_dimensions["A"].width = 2
ws.column_dimensions["B"].width = 26
ws.column_dimensions["C"].width = 62
ws.column_dimensions["D"].width = 40
ws["B1"] = "Business Data Workbook — Start Here"
ws["B1"].font = f_title
ws["B2"] = "One place for every client, job, hour and dollar — so decisions about pricing, what work to chase, and hiring come from data, not gut feel."
ws["B2"].font = f_sub
r = 4

def para(text, bold=False, h=None, col=2, span=3):
    global r
    ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + span - 1)
    c = ws.cell(r, col, text)
    c.font = f_bold if bold else f_body
    c.alignment = wrap_l
    ws.row_dimensions[r].height = h or max(15, 15 * (len(text) // 120 + 1))
    r += 1

def table(headers, rows):
    global r
    for j, h in enumerate(headers):
        hdr(ws, r, 2 + j, h)
    r += 1
    for row in rows:
        mx = 1
        for j, v in enumerate(row):
            c = ws.cell(r, 2 + j, v)
            c.font = f_bold if j == 0 else f_body
            c.alignment, c.border = wrap_l, box
            width = [26, 62, 40][j]
            mx = max(mx, len(str(v)) // int(width * 1.1) + 1)
        ws.row_dimensions[r].height = 14 * mx + 2
        r += 1
    r += 1

section(ws, r, "FIRST-TIME SETUP (about 30 minutes)", 2, 4); r += 1
for t in [
    "1.  Settings tab → set the yellow cells: business name, labor burden %, target gross margin %, target net profit %. Review the dropdown lists (trades, project types, lead sources…) and add or rename anything that doesn't fit how you work.",
    "2.  Team tab → add yourself, your partners and anyone else who works jobs. Set a Base Pay $/hr for EVERY person — including owners (see rule #2 below). The two rows there now hold placeholder rates.",
    "3.  Pricing & Hiring tab → fill in the monthly overhead budget (insurance, truck, phone, software, marketing…) and your capacity (number of field workers, hours per year).",
    "4.  Delete the example rows (yellow-highlighted notes say 'EXAMPLE') on each data tab: right-click the row number → Delete. Deleting whole rows keeps all formulas working.",
    "5.  Enter your active and recent jobs. Even 10–20 past jobs with rough hours and costs will start showing patterns.",
]:
    para(t)
r += 1

if GSHEETS:
    section(ws, r, "LOGGING FROM WHATSAPP", 2, 4); r += 1
    for t in [
        "Once the WhatsApp bot is connected (see the setup guide in the whatsapp-bot folder), text the business WhatsApp number in plain English and it fills in the right tabs, then replies with exactly what it wrote. Only phone numbers listed on the Team tab can write.",
        "Examples:  'Me and Sam 8 to 4:30 on the Johnson AC job, install, 30 min lunch'  ·  a photo of a receipt with 'Johnson job'  ·  'New lead: Maria Lopez 312-555-0199, kitchen remodel at 44 Oak St, found us on Google'  ·  'Got paid 4,600 by check from Johnson'  ·  'Johnson job done today'  ·  'Q-0012 lost, went with ABC at 9,500'.",
        "Reply UNDO to reverse the last thing the bot wrote. Send STATS for this month's numbers. Every message and every cell written is recorded on the WhatsApp Log tab.",
    ]:
        para(t)
    r += 1

section(ws, r, "HOW TO READ THE TABS", 2, 4); r += 1
para("Blue header / blue text = you type it.   Gray header / gray cell = calculates automatically — don't type over it.   Yellow cell = a key assumption you should set.   Hover over any header with a red corner (or a note) for an explanation.")
para("Tab colors:  Blue tabs = data you enter.  Green tabs = analysis (all automatic).  Gold = Dashboard.  Gray = Settings." + ("  WhatsApp green = the bot's log." if GSHEETS else ""))
r += 1
table(["Tab", "What goes in it", "When to update"], [
    ("Dashboard", "Headline numbers + trade scoreboard + things that need attention. Pick the analysis period here (All Time or a year).", "Look at it weekly"),
    ("Clients", "Every client: contact info, type, how they found you. Auto: lifetime revenue/profit, repeat status, tier.", "New client"),
    ("Properties", "Every job site: address, distance, building details, equipment ages. Auto: future-work flags (old AC, old water heater, small panel).", "First visit to a site"),
    ("Quotes", "Every estimate, won or lost, with estimated hours and costs, competitor price, and lost reason.", "Every estimate you give"),
    ("Projects", "Every job. You fill the basics and dates; profit, hours, variance, collections calculate automatically.", "Job booked, started, finished"),
    ("Change Orders", "Every scope/price change after agreement. Only 'Approved' ones count.", "When scope changes"),
    ("Time Log", "Every paid hour: who, which job, which task. THE most important tab.", "Daily (5 minutes at end of day)"),
    ("Expenses", "Every receipt/bill: job cost (Project ID) or overhead (OVERHEAD).", "Daily or weekly"),
    ("Invoices", "Every invoice and payment.", "When you bill / get paid"),
    ("Team", "Everyone who works jobs, their pay and burden. Auto: hours, wrench-time %, revenue generated, value per $1 of pay.", "Hires, raises"),
    ("Vendors", "Suppliers & subs: terms, discounts, insurance certificate expiry. Auto: spend.", "New vendor"),
    ("Callbacks", "Every return trip to fix finished work, with root cause.", "When it happens"),
    ("Market Rates", "Competitor prices you learn (their quote, lost-job feedback, websites, calls).", "Whenever you learn one"),
    ("Trade Analysis", "Revenue, profit, profit per hour, estimate accuracy, win rate and market comparison for each trade — with a verdict.", "Automatic"),
    ("Project Types", "Same, for every kind of job (AC replacement, LVP install, bathroom remodel…). Where you decide what to chase or reprice.", "Automatic"),
    ("Time Analysis", "Hours by trade × task and by person × task. Shows where you're slow.", "Automatic"),
    ("Pricing & Hiring", "Overhead budget, break-even and target hourly rates, rate by trade vs market, and a hiring calculator.", "Overhead budget: yearly"),
    ("Marketing", "Leads, win rate, revenue, profit, spend, cost per lead and ROI for each lead source.", "Automatic"),
    ("Clients & Areas", "Profit by client type and travel zone, client tiers, top 10 clients.", "Automatic"),
    ("Monthly Trend", "Month-by-month revenue, profit, overhead, cash, hours, quotes for the chosen year, with chart.", "Automatic (set the year)"),
    ("Settings", "Assumptions and every dropdown list.", "Rarely"),
] + ([("WhatsApp Log", "Every WhatsApp message the bot handled and exactly which cells it wrote.", "Automatic")] if GSHEETS else []))

section(ws, r, "THE 8 RULES THAT MAKE THE DATA TRUSTWORTHY", 2, 4); r += 1
for t in [
    "1.  Log EVERY paid hour on Time Log — travel, material runs, estimates, admin, callbacks. Time you don't log is time you can't price. Non-job time goes to Project ID 'OVERHEAD'.",
    "2.  Owners get a pay rate. On Team, set your Base Pay to what you'd have to pay someone to replace you on the tools. If owners are $0, every job looks more profitable than it is, and you'll never know if a hire can take over that work profitably.",
    "3.  Every receipt goes on Expenses with a Project ID or OVERHEAD. Materials bought for a job go to that job — even if you return some later (log the return as a negative amount).",
    "4.  Every quote goes on Quotes — especially lost ones. Ask why you lost and what the other company charged. That is free market research.",
    "5.  Close out jobs: set Status = Completed AND fill the Actual Finish date. The analysis tabs only count completed jobs (in-progress jobs have partial costs and would mislead you).",
    "6.  IDs never change. C-0001, P-0001, Q-0001, J-0001, W-01… Use the 'Next ID' hint at the top of each tab. The same Project ID must appear on Time Log, Expenses, Invoices and Change Orders.",
    "7.  Pick Project Types consistently. The analysis groups by them. If a job doesn't fit, add a new type on Settings rather than forcing it.",
    "8.  Check the red 'Data Check' column on Time Log and Expenses, and the 'Needs Attention' box on the Dashboard, once a week.",
]:
    para(t)
r += 1

section(ws, r, "KEY NUMBERS — WHAT THEY MEAN", 2, 4); r += 1
table(["Metric", "Meaning", "Why it matters"], [
    ("Gross Profit / Labor Hour", "(Revenue − direct job costs) ÷ crew hours on the job.", "THE number for choosing work. A $20k job that ties up 300 hours can be worse than a $3k job done in 15."),
    ("Gross Margin %", "Gross profit ÷ revenue.", "Tells you if the price covers the job. Compare with your target on Settings."),
    ("Loaded Labor Cost", "Hourly pay × (1 + burden %). Burden = payroll taxes, workers' comp, benefits.", "An employee paid $25/hr can cost you $31+/hr."),
    ("Overhead per Job Hour", "All overhead (insurance, truck, phone, marketing, non-job labor…) ÷ hours billed to jobs.", "Every job hour must earn this before you make any profit."),
    ("Break-Even Rate", "Loaded labor cost + overhead per job hour.", "Charging less than this per labor hour loses money."),
    ("Target Rate", "Break-even ÷ (1 − target net profit %).", "Your minimum labor rate. Compare with the market on Pricing & Hiring."),
    ("Wrench-Time %", "Hours spent actually doing the work ÷ all paid hours.", "The gap is travel, pickups, waiting, admin — the easiest place to find more capacity without hiring."),
    ("Hours vs Estimate %", "Actual hours ÷ estimated hours − 1.", "Consistently + for a job type = you're under-quoting it."),
    ("Win Rate", "Won ÷ (won + lost + expired).", "Very high win rate + high profit/hr = you're probably too cheap. Low win rate = price, speed or trust issue."),
    ("Revenue per $1 of Labor", "Revenue generated by a person's hours ÷ what those hours cost.", "Compares what each team member is worth to the business."),
    ("Marketing ROI", "(Gross profit from a lead source − its marketing spend) ÷ spend.", "Tells you which ads to scale and which to cut."),
    ("Days to Get Paid", "Days from finishing a job to being paid in full.", "Slow payers tie up your cash; tighten deposits and terms."),
])

section(ws, r, "QUESTIONS THIS WORKBOOK ANSWERS", 2, 4); r += 1
table(["Question", "Where to look", ""], [
    ("Which jobs make the most money for the least time?", "Project Types → sort by Gross Profit / Labor Hr, read the Verdict column", ""),
    ("What should we charge per hour?", "Pricing & Hiring → sections 3 and 4 (break-even, target, market, suggested)", ""),
    ("Are we cheaper or pricier than competitors?", "Trade Analysis / Project Types → market columns; Quotes → competitor price", ""),
    ("Where are we slowest?", "Time Analysis; Project Types → Hours vs Estimate; Dashboard → travel/pickup/wait hours", ""),
    ("Can we afford to hire? What would they need to produce?", "Pricing & Hiring → section 5 (hiring calculator)", ""),
    ("What is each person worth to the business?", "Team → revenue generated, revenue per $1 of labor, wrench-time %", ""),
    ("Which marketing works?", "Marketing tab", ""),
    ("Who are our best clients? Who should we call?", "Clients & Areas → tiers and top 10; Clients → days since last project; Properties → future-work flags", ""),
    ("Are far-away jobs worth it?", "Clients & Areas → by travel zone", ""),
    ("Are we growing? Is cash coming in?", "Monthly Trend; Dashboard", ""),
])

section(ws, r, "HOUSEKEEPING", 2, 4); r += 1
for t in [
    "• Formulas are pre-filled for: Clients 500 rows, Properties 500, Quotes 1,000, Projects 500, Change Orders 300, Time Log 3,000, Expenses 3,000, Invoices 1,000, Team 30, Vendors 150, Callbacks 200, Market Rates 500. When you get close, select the last filled row and drag its fill handle down — every analysis uses whole columns, so new rows are picked up automatically.",
    "• Add new columns at the far right of a tab, not in the middle.",
    "• Use the filter arrows in each header row to sort and search. Sorting a data tab is safe.",
    "• Revenue is counted when a project is Completed (in the month of its Actual Finish date). Cash timing is tracked separately on Invoices.",
    ("• Google Sheets keeps full version history (File → Version history), so you can always roll back." if GSHEETS else
     "• Keep a backup: save a dated copy monthly (e.g. Business_Data_2026-09.xlsx) or keep the file in OneDrive / Google Drive."),
    "• This workbook is built so its tabs can later be moved into a database or job-management app (Jobber, Housecall Pro, ServiceTitan, QuickBooks) without losing your history — the IDs and column names are the structure.",
]:
    para(t)

# ================================================================ Google Sheets compatibility pass
if GSHEETS:
    def to_indirect(formula):
        def rep(m):
            nm = m.group(0)
            ref = NAMES[nm]
            return f'INDIRECT("{ref}")'
        keys = sorted([k for k, v in NAMES.items() if not v.startswith("OFFSET")], key=len, reverse=True)
        pat = re.compile(r"(?<![A-Za-z_!$\"])(" + "|".join(map(re.escape, keys)) + r")(?![A-Za-z0-9_(])")
        return pat.sub(rep, formula)

    for ws in wb.worksheets:
        # dropdowns -> plain ranges
        for v in ws.data_validations.dataValidation:
            f1 = (v.formula1 or "").lstrip("=")
            if f1 in NAMES:
                v.formula1 = NAMES[f1]
        # conditional formatting: Sheets can't see named ranges / other sheets here -> INDIRECT
        for cfr in ws.conditional_formatting:
            for rule in cfr.rules:
                if rule.formula:
                    rule.formula = [to_indirect(f) for f in rule.formula]
        # _xlfn prefixes are an Excel storage detail
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "_xlfn." in c.value:
                    c.value = c.value.replace("_xlfn.", "")

wb.calculation.fullCalcOnLoad = True
wb.active = 0
wb.save(OUT)
print("saved", OUT)
